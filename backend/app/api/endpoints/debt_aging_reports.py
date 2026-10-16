"""
Endpoint Báo cáo Công nợ Quá hạn theo Tuổi nợ (Story S8-09 - SCRUM-280).

API:
- GET /api/v1/debt-aging-reports        → Lấy báo cáo tuổi nợ (JSON)
- GET /api/v1/debt-aging-reports/export → Xuất Excel (.xlsx)

Phân quyền: Director, SalesManager, Accountant, Admin được xem toàn bộ.
             SalesStaff chỉ xem đại lý mình phụ trách.

Tuân thủ AGENTS.md - Additive Only: File mới hoàn toàn, không sửa file cũ.
"""
from typing import Optional
from datetime import date
import io

from fastapi import APIRouter, Depends, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import lay_phien_db
from app.core.dependencies import lay_nguoi_dung_hien_tai
from app.models.auth import User
from app.schemas.debt_aging_report import DebtAgingReportResponse
from app.services.debt_aging_report_service import DebtAgingReportService

router = APIRouter(
    prefix="/debt-aging-reports",
    tags=["Báo cáo Công nợ Quá hạn theo Tuổi nợ (S8-09)"],
)


@router.get(
    "",
    response_model=DebtAgingReportResponse,
    status_code=status.HTTP_200_OK,
    summary="Xem báo cáo công nợ quá hạn theo tuổi nợ (S8-09 - SCRUM-280)",
)
def bao_cao_cong_no_qua_han(
    customer_id: Optional[str] = Query(None, description="Lọc theo mã đại lý"),
    sales_rep: Optional[str] = Query(None, description="Lọc theo tên nhân viên phụ trách"),
    region: Optional[str] = Query(None, description="Lọc theo khu vực"),
    as_of_date: Optional[date] = Query(None, description="Ngày lập báo cáo (mặc định hôm nay, định dạng YYYY-MM-DD)"),
    current_user: User = Depends(lay_nguoi_dung_hien_tai),
    db: Session = Depends(lay_phien_db),
):
    """
    Báo cáo Công nợ Quá hạn theo Tuổi nợ cho Kế toán công nợ (AC-1, AC-2, AC-3):

    - **AC-1 – Tổng hợp theo nhóm tuổi nợ**: Phân loại dư nợ theo 5 nhóm:
      Chưa đến hạn / 1-30 ngày / 31-60 ngày / 61-90 ngày / Trên 90 ngày.
    - **AC-2 – Tổng hợp theo nhân viên phụ trách**: Thống kê dư nợ và tỷ lệ quá hạn theo từng NVKD.
    - **AC-3 – Xác định đại lý rủi ro cao nhất**: Trả về `top_risk_customers` sắp xếp theo mức độ
      rủi ro (CRITICAL → HIGH → MEDIUM → LOW) và số tiền quá hạn giảm dần.

    Phân quyền: Director / SalesManager / Accountant xem tất cả. SalesStaff chỉ xem đại lý mình phụ trách.
    """
    return DebtAgingReportService.get_debt_aging_report(
        db=db,
        current_user=current_user,
        customer_id=customer_id,
        sales_rep=sales_rep,
        region=region,
        as_of_date=as_of_date,
    )


@router.get(
    "/export",
    status_code=status.HTTP_200_OK,
    summary="Xuất Excel báo cáo công nợ quá hạn theo tuổi nợ (S8-09 - AC-4 Xuất Excel)",
)
def xuat_excel_cong_no_qua_han(
    customer_id: Optional[str] = Query(None, description="Lọc theo mã đại lý"),
    sales_rep: Optional[str] = Query(None, description="Lọc theo tên nhân viên phụ trách"),
    region: Optional[str] = Query(None, description="Lọc theo khu vực"),
    as_of_date: Optional[date] = Query(None, description="Ngày lập báo cáo (mặc định hôm nay)"),
    current_user: User = Depends(lay_nguoi_dung_hien_tai),
    db: Session = Depends(lay_phien_db),
):
    """
    Xuất file Excel (.xlsx) báo cáo công nợ quá hạn theo tuổi nợ (AC-4).

    File Excel gồm 3 sheet:
    - **Sheet 1 – Chi tiết đại lý**: Danh sách đại lý và phân tích tuổi nợ từng đại lý.
    - **Sheet 2 – Tổng hợp theo NVKD**: Dư nợ và tỷ lệ quá hạn theo từng nhân viên phụ trách.
    - **Sheet 3 – Top rủi ro**: 10 đại lý rủi ro cao nhất cần ưu tiên xử lý.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(
            status_code=status.HTTP_501_NOT_IMPLEMENTED,
            detail="Thư viện openpyxl chưa được cài đặt. Chạy: pip install openpyxl",
        )

    report = DebtAgingReportService.get_debt_aging_report(
        db=db,
        current_user=current_user,
        customer_id=customer_id,
        sales_rep=sales_rep,
        region=region,
        as_of_date=as_of_date,
    )

    wb = openpyxl.Workbook()

    # ===== STYLE HELPERS =====
    def _header_style(cell, color: str = "1F4E79"):
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _risk_color(risk: str) -> str:
        return {"CRITICAL": "C00000", "HIGH": "FF0000", "MEDIUM": "FFC000", "LOW": "70AD47"}.get(risk, "FFFFFF")

    def _money(val: int) -> str:
        return f"{val:,}"

    # ===== SHEET 1: Chi tiết đại lý =====
    ws1 = wb.active
    ws1.title = "Chi tiết đại lý"

    # Tiêu đề lớn
    ws1.merge_cells("A1:M1")
    ws1["A1"] = f"BÁO CÁO CÔNG NỢ QUÁ HẠN THEO TUỔI NỢ – Ngày lập: {report.report_date}"
    ws1["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws1["A1"].alignment = Alignment(horizontal="center")

    # Tóm tắt
    ws1["A2"] = f"Tổng đại lý có nợ: {report.total_customers_with_debt}"
    ws1["D2"] = f"Tổng dư nợ: {_money(report.total_debt_amount)} VNĐ"
    ws1["H2"] = f"Tổng quá hạn: {_money(report.total_overdue_amount)} VNĐ"

    # Header bảng
    headers_s1 = [
        "STT", "Mã đại lý", "Tên đại lý", "Khu vực", "NVKD phụ trách",
        "Tổng dư nợ (VNĐ)", "Chưa đến hạn", "1-30 ngày", "31-60 ngày", "61-90 ngày",
        "Trên 90 ngày", "Tổng quá hạn (VNĐ)", "Mức rủi ro",
    ]
    for col, h in enumerate(headers_s1, 1):
        cell = ws1.cell(row=4, column=col, value=h)
        _header_style(cell)

    for idx, row in enumerate(report.customers, 1):
        r = idx + 4
        ws1.cell(row=r, column=1, value=idx)
        ws1.cell(row=r, column=2, value=row.customer_id)
        ws1.cell(row=r, column=3, value=row.customer_name)
        ws1.cell(row=r, column=4, value=row.region or "")
        ws1.cell(row=r, column=5, value=row.assigned_sales_rep or "Chưa phân công")
        ws1.cell(row=r, column=6, value=row.total_debt)
        ws1.cell(row=r, column=7, value=row.current_amount)
        ws1.cell(row=r, column=8, value=row.bucket_1_30)
        ws1.cell(row=r, column=9, value=row.bucket_31_60)
        ws1.cell(row=r, column=10, value=row.bucket_61_90)
        ws1.cell(row=r, column=11, value=row.bucket_over_90)
        ws1.cell(row=r, column=12, value=row.overdue_total)
        risk_cell = ws1.cell(row=r, column=13, value=row.risk_level)
        risk_cell.fill = PatternFill("solid", fgColor=_risk_color(row.risk_level))
        risk_cell.font = Font(bold=True, color="FFFFFF")
        risk_cell.alignment = Alignment(horizontal="center")

    # Dòng tổng cộng
    total_row = len(report.customers) + 5
    ws1.cell(row=total_row, column=3, value="TỔNG CỘNG").font = Font(bold=True)
    ws1.cell(row=total_row, column=6, value=report.total_debt_amount).font = Font(bold=True)
    ws1.cell(row=total_row, column=12, value=report.total_overdue_amount).font = Font(bold=True)

    # Độ rộng cột
    col_widths_s1 = [5, 14, 30, 15, 20, 18, 16, 14, 14, 14, 16, 18, 12]
    for i, w in enumerate(col_widths_s1, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ===== SHEET 2: Tổng hợp theo NVKD =====
    ws2 = wb.create_sheet("Tổng hợp theo NVKD")
    ws2.merge_cells("A1:J1")
    ws2["A1"] = f"TỔNG HỢP CÔNG NỢ THEO NHÂN VIÊN PHỤ TRÁCH – Ngày: {report.report_date}"
    ws2["A1"].font = Font(bold=True, size=12, color="1F4E79")
    ws2["A1"].alignment = Alignment(horizontal="center")

    headers_s2 = [
        "Nhân viên phụ trách", "Số đại lý", "Tổng dư nợ (VNĐ)", "Chưa đến hạn",
        "1-30 ngày", "31-60 ngày", "61-90 ngày", "Trên 90 ngày",
        "Tổng quá hạn (VNĐ)", "Tỷ lệ quá hạn (%)", "ĐL rủi ro cao",
    ]
    for col, h in enumerate(headers_s2, 1):
        cell = ws2.cell(row=3, column=col, value=h)
        _header_style(cell, "2E75B6")

    for idx, rep in enumerate(report.by_sales_rep, 1):
        r = idx + 3
        ws2.cell(row=r, column=1, value=rep.sales_rep_name)
        ws2.cell(row=r, column=2, value=rep.customer_count)
        ws2.cell(row=r, column=3, value=rep.total_debt)
        ws2.cell(row=r, column=4, value=rep.current_amount)
        ws2.cell(row=r, column=5, value=rep.bucket_1_30)
        ws2.cell(row=r, column=6, value=rep.bucket_31_60)
        ws2.cell(row=r, column=7, value=rep.bucket_61_90)
        ws2.cell(row=r, column=8, value=rep.bucket_over_90)
        ws2.cell(row=r, column=9, value=rep.overdue_total)
        ws2.cell(row=r, column=10, value=round(rep.overdue_ratio * 100, 2))
        ws2.cell(row=r, column=11, value=rep.high_risk_customers)

    for i, w in enumerate([25, 12, 18, 16, 14, 14, 14, 16, 18, 16, 14], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ===== SHEET 3: Top đại lý rủi ro cao nhất =====
    ws3 = wb.create_sheet("Top rủi ro")
    ws3.merge_cells("A1:I1")
    ws3["A1"] = f"TOP ĐẠI LÝ RỦI RO CÔNG NỢ CAO NHẤT – Ngày: {report.report_date}"
    ws3["A1"].font = Font(bold=True, size=12, color="C00000")
    ws3["A1"].alignment = Alignment(horizontal="center")

    headers_s3 = [
        "STT", "Tên đại lý", "NVKD phụ trách", "Tổng dư nợ (VNĐ)",
        "Tổng quá hạn (VNĐ)", "Tỷ lệ quá hạn (%)",
        "Quá hạn lâu nhất (ngày)", "Hóa đơn cũ nhất", "Mức rủi ro",
    ]
    for col, h in enumerate(headers_s3, 1):
        cell = ws3.cell(row=3, column=col, value=h)
        _header_style(cell, "C00000")

    for idx, row in enumerate(report.top_risk_customers, 1):
        r = idx + 3
        ws3.cell(row=r, column=1, value=idx)
        ws3.cell(row=r, column=2, value=row.customer_name)
        ws3.cell(row=r, column=3, value=row.assigned_sales_rep or "Chưa phân công")
        ws3.cell(row=r, column=4, value=row.total_debt)
        ws3.cell(row=r, column=5, value=row.overdue_total)
        ws3.cell(row=r, column=6, value=round(row.overdue_ratio * 100, 2))
        ws3.cell(row=r, column=7, value=row.oldest_overdue_days)
        ws3.cell(row=r, column=8, value=row.oldest_invoice_id or "")
        risk_cell = ws3.cell(row=r, column=9, value=row.risk_level)
        risk_cell.fill = PatternFill("solid", fgColor=_risk_color(row.risk_level))
        risk_cell.font = Font(bold=True, color="FFFFFF")
        risk_cell.alignment = Alignment(horizontal="center")

    for i, w in enumerate([5, 30, 25, 18, 18, 16, 22, 18, 12], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # ===== Xuất file =====
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"BaoCaoCongNoQuaHan_{report.report_date.replace('/', '-')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )
