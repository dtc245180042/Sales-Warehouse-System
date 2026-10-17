import io
from typing import List, Optional, Tuple, Dict, Any
from datetime import datetime, date, timezone
from zoneinfo import ZoneInfo
import logging

from sqlalchemy.orm import Session
from sqlalchemy import or_, and_, func
from fastapi import HTTPException, status

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.auth import User, UserRole
from app.models.customer import Customer
from app.models.customer_credit_profile import CustomerCreditProfile
from app.models.customer_debt_profile import CustomerDebtProfile
from app.models.invoice import Invoice, InvoiceStatus
from app.schemas.agency_debt_report import (
    AgencyDebtReportItem,
    AgencyDebtSummary,
    AgencyDebtReportResponse,
    AgencyInvoiceDetail,
    AgencyDebtDetailResponse,
)

logger = logging.getLogger(__name__)
TZ_VN = ZoneInfo("Asia/Ho_Chi_Minh")


class AgencyDebtService:
    """
    Service xử lý toàn bộ nghiệp vụ Sổ công nợ theo đại lý & Phân tích tuổi nợ (Story S7-03).
    Bao gồm:
    - SCRUM-717: Tính nhóm tuổi nợ (0-30, 31-60, 61-90, >90) và xác định khoản quá hạn theo due_date.
    - SCRUM-718: Tổng hợp các chỉ tiêu phải thu, đã thu, còn nợ theo đại lý.
    - SCRUM-719: Bộ lọc linh hoạt (search, region, customer_group, only_overdue, aging_group) & tối ưu Zero N+1.
    - SCRUM-721: Tính toán dòng tổng cộng ở chân bảng.
    - SCRUM-722: Phát triển chức năng xuất Excel chuẩn kế toán.
    """

    @staticmethod
    def calculate_invoice_aging(invoice: Invoice, as_of_date: date) -> Dict[str, Any]:
        """
        Tính tuổi nợ và tình trạng quá hạn của một hóa đơn cụ thể (SCRUM-717).
        - Ngày công nợ: ngày phát hành hóa đơn (invoice.issued_at)
        - Tuổi nợ (aging_days): số ngày từ ngày phát hành đến ngày chốt báo cáo (as_of_date)
        - Nhóm tuổi nợ: 0-30, 31-60, 61-90, >90 ngày
        - Khoản quá hạn: nếu as_of_date > invoice.due_date và còn nợ > 0
        """
        if isinstance(invoice.issued_at, datetime):
            issued_date = invoice.issued_at.date()
        else:
            issued_date = invoice.issued_at

        if isinstance(invoice.due_date, datetime):
            due_date = invoice.due_date.date()
        else:
            due_date = invoice.due_date

        aging_days = max(0, (as_of_date - issued_date).days)
        
        # Phân chia 4 nhóm tuổi nợ chuẩn theo AC
        if aging_days <= 30:
            aging_bucket = "0_30"
        elif aging_days <= 60:
            aging_bucket = "31_60"
        elif aging_days <= 90:
            aging_bucket = "61_90"
        else:
            aging_bucket = "over_90"

        # Tính quá hạn dựa trên ngày đến hạn (due_date)
        days_overdue = (as_of_date - due_date).days
        is_overdue = bool(days_overdue > 0 and invoice.remaining_amount > 0)
        overdue_amount = invoice.remaining_amount if is_overdue else 0
        normalized_days_overdue = max(0, days_overdue) if is_overdue else 0

        return {
            "aging_days": aging_days,
            "aging_bucket": aging_bucket,
            "days_overdue": normalized_days_overdue,
            "is_overdue": is_overdue,
            "overdue_amount": overdue_amount,
        }

    @classmethod
    def get_agency_debt_data(
        cls,
        db: Session,
        current_user: User,
        as_of_date: Optional[date] = None,
        search: Optional[str] = None,
        region: Optional[str] = None,
        customer_group: Optional[str] = None,
        assigned_sales_rep: Optional[str] = None,
        only_overdue: bool = False,
        aging_group: Optional[str] = None,
        has_debt_only: bool = True,
        customer_id: Optional[str] = None,
    ) -> Tuple[List[AgencyDebtReportItem], AgencyDebtSummary, date]:
        """
        Lấy toàn bộ dữ liệu báo cáo sổ công nợ đại lý đã tổng hợp, tối ưu Zero N+1 query.
        """
        if as_of_date is None:
            as_of_date = datetime.now(TZ_VN).date()

        # 1. Truy vấn danh sách đại lý (Customers) có áp dụng phân quyền và bộ lọc
        cust_query = db.query(Customer)

        # Phân quyền Customer Scope (anti-IDOR): Đại lý chỉ xem được của chính mình
        user_role = getattr(current_user, "role", "")
        if isinstance(user_role, UserRole):
            user_role = user_role.value

        if user_role == UserRole.CUSTOMER.value:
            user_cust_id = getattr(current_user, "customer_id", None)
            if user_cust_id:
                cust_query = cust_query.filter(Customer.id == user_cust_id)
            else:
                user_cust = db.query(Customer).filter(
                    or_(Customer.email == current_user.email, Customer.phone == current_user.phone_number)
                ).first()
                if user_cust:
                    cust_query = cust_query.filter(Customer.id == user_cust.id)
                else:
                    cust_query = cust_query.filter(Customer.id == "CUSTOMER_NOT_FOUND")
        elif user_role == UserRole.SALES_REP.value:
            # NVKD: Nếu được gán phụ trách đại lý nào thì ưu tiên lọc theo đại lý đó
            if current_user.username:
                cust_query = cust_query.filter(
                    or_(
                        Customer.assigned_sales_rep == current_user.username,
                        Customer.assigned_sales_rep == current_user.full_name,
                        Customer.assigned_sales_rep.is_(None)
                    )
                )

        if customer_id:
            cust_query = cust_query.filter(Customer.id == customer_id)

        if search:
            search_pattern = f"%{search.strip()}%"
            cust_query = cust_query.filter(
                or_(
                    Customer.code.ilike(search_pattern),
                    Customer.name.ilike(search_pattern),
                    Customer.phone.ilike(search_pattern),
                )
            )

        if region:
            cust_query = cust_query.filter(Customer.region == region)

        if customer_group:
            cust_query = cust_query.filter(Customer.customer_group == customer_group)

        if assigned_sales_rep:
            cust_query = cust_query.filter(Customer.assigned_sales_rep == assigned_sales_rep)

        customers = cust_query.order_by(Customer.code.asc()).all()
        if not customers:
            return [], AgencyDebtSummary(), as_of_date

        customer_ids = [c.id for c in customers]

        # 2. Tối ưu Zero N+1: Truy vấn toàn bộ Credit Profile trong 1 câu lệnh duy nhất
        credit_profiles = db.query(CustomerCreditProfile).filter(
            CustomerCreditProfile.customer_id.in_(customer_ids)
        ).all()
        credit_map: Dict[str, CustomerCreditProfile] = {cp.customer_id: cp for cp in credit_profiles}

        # Dự phòng CustomerDebtProfile nếu có
        debt_profiles = db.query(CustomerDebtProfile).filter(
            CustomerDebtProfile.customer_id.in_(customer_ids)
        ).all()
        debt_map: Dict[str, CustomerDebtProfile] = {dp.customer_id: dp for dp in debt_profiles}

        # 3. Tối ưu Zero N+1: Truy vấn toàn bộ hóa đơn liên quan trong 1 câu lệnh duy nhất
        invoices = db.query(Invoice).filter(
            Invoice.customer_id.in_(customer_ids)
        ).order_by(Invoice.issued_at.asc()).all()

        invoices_by_customer: Dict[str, List[Invoice]] = {cid: [] for cid in customer_ids}
        for inv in invoices:
            invoices_by_customer[inv.customer_id].append(inv)

        # 4. Tính toán và tổng hợp từng đại lý
        report_items: List[AgencyDebtReportItem] = []

        for cust in customers:
            cust_invs = invoices_by_customer.get(cust.id, [])
            cp = credit_map.get(cust.id)
            dp = debt_map.get(cust.id)

            credit_limit = 0
            max_debt_days = 0
            if cp:
                credit_limit = int(cp.credit_limit or 0)
                max_debt_days = int(cp.max_debt_days or 0)
            elif dp:
                credit_limit = int(dp.credit_limit or 0)

            total_receivable = sum(int(inv.total_amount or 0) for inv in cust_invs)
            total_paid = sum(int(inv.paid_amount or 0) for inv in cust_invs)
            total_debt = sum(int(inv.remaining_amount or 0) for inv in cust_invs)

            debt_0_30 = 0
            debt_31_60 = 0
            debt_61_90 = 0
            debt_over_90 = 0
            overdue_debt = 0
            max_overdue_days = 0
            is_overdue = False
            unpaid_count = 0
            oldest_date = None

            for inv in cust_invs:
                rem = int(inv.remaining_amount or 0)
                if rem > 0:
                    unpaid_count += 1
                    if oldest_date is None or inv.issued_at < oldest_date:
                        oldest_date = inv.issued_at

                    aging_info = cls.calculate_invoice_aging(inv, as_of_date)
                    bucket = aging_info["aging_bucket"]
                    if bucket == "0_30":
                        debt_0_30 += rem
                    elif bucket == "31_60":
                        debt_31_60 += rem
                    elif bucket == "61_90":
                        debt_61_90 += rem
                    elif bucket == "over_90":
                        debt_over_90 += rem

                    if aging_info["is_overdue"]:
                        is_overdue = True
                        overdue_debt += aging_info["overdue_amount"]
                        if aging_info["days_overdue"] > max_overdue_days:
                            max_overdue_days = aging_info["days_overdue"]

            item = AgencyDebtReportItem(
                customer_id=cust.id,
                customer_code=cust.code,
                customer_name=cust.name,
                phone=cust.phone,
                address=cust.address,
                tax_code=cust.tax_code,
                region=cust.region,
                customer_group=cust.customer_group,
                assigned_sales_rep=cust.assigned_sales_rep,
                credit_limit=credit_limit,
                max_debt_days=max_debt_days,
                total_receivable=total_receivable,
                total_paid=total_paid,
                total_debt=total_debt,
                debt_0_30=debt_0_30,
                debt_31_60=debt_31_60,
                debt_61_90=debt_61_90,
                debt_over_90=debt_over_90,
                overdue_debt=overdue_debt,
                max_overdue_days=max_overdue_days,
                is_overdue=is_overdue,
                unpaid_invoices_count=unpaid_count,
                oldest_invoice_date=oldest_date,
            )
            report_items.append(item)

        # 5. Áp dụng các bộ lọc chuyên sâu (SCRUM-719)
        filtered_items = []
        for it in report_items:
            # Lọc chỉ hiện đại lý có công nợ hoặc có phát sinh hóa đơn
            if has_debt_only and it.total_debt <= 0 and it.total_receivable <= 0:
                continue

            # Lọc riêng nhóm quá hạn (Tiêu chí AC)
            if only_overdue and not it.is_overdue:
                continue

            # Lọc theo nhóm tuổi nợ cụ thể
            if aging_group:
                if aging_group == "0_30" and it.debt_0_30 <= 0:
                    continue
                elif aging_group == "31_60" and it.debt_31_60 <= 0:
                    continue
                elif aging_group == "61_90" and it.debt_61_90 <= 0:
                    continue
                elif aging_group in ("over_90", ">90") and it.debt_over_90 <= 0:
                    continue

            filtered_items.append(it)

        # 6. Tính toán dòng TỔNG CỘNG ở chân bảng (SCRUM-721)
        summary = AgencyDebtSummary(
            total_customers=len(filtered_items),
            total_receivable=sum(x.total_receivable for x in filtered_items),
            total_paid=sum(x.total_paid for x in filtered_items),
            total_debt=sum(x.total_debt for x in filtered_items),
            total_debt_0_30=sum(x.debt_0_30 for x in filtered_items),
            total_debt_31_60=sum(x.debt_31_60 for x in filtered_items),
            total_debt_61_90=sum(x.debt_61_90 for x in filtered_items),
            total_debt_over_90=sum(x.debt_over_90 for x in filtered_items),
            total_overdue_debt=sum(x.overdue_debt for x in filtered_items),
            overdue_customers_count=sum(1 for x in filtered_items if x.is_overdue),
        )

        return filtered_items, summary, as_of_date

    @classmethod
    def get_agency_debt_report(
        cls,
        db: Session,
        current_user: User,
        as_of_date: Optional[date] = None,
        search: Optional[str] = None,
        region: Optional[str] = None,
        customer_group: Optional[str] = None,
        assigned_sales_rep: Optional[str] = None,
        only_overdue: bool = False,
        aging_group: Optional[str] = None,
        has_debt_only: bool = True,
        page: int = 1,
        limit: int = 50,
    ) -> AgencyDebtReportResponse:
        """
        API Lấy danh sách sổ công nợ theo đại lý có phân trang và dòng tổng cộng chân bảng (SCRUM-718 & 721).
        """
        all_items, summary, calculated_date = cls.get_agency_debt_data(
            db=db,
            current_user=current_user,
            as_of_date=as_of_date,
            search=search,
            region=region,
            customer_group=customer_group,
            assigned_sales_rep=assigned_sales_rep,
            only_overdue=only_overdue,
            aging_group=aging_group,
            has_debt_only=has_debt_only,
        )

        total_count = len(all_items)
        start_idx = (page - 1) * limit
        end_idx = start_idx + limit
        paginated_items = all_items[start_idx:end_idx]

        return AgencyDebtReportResponse(
            items=paginated_items,
            summary=summary,
            total=total_count,
            page=page,
            limit=limit,
            as_of_date=calculated_date.strftime("%Y-%m-%d"),
        )

    @classmethod
    def get_agency_invoices_detail(
        cls,
        db: Session,
        customer_id: str,
        current_user: User,
        as_of_date: Optional[date] = None,
    ) -> AgencyDebtDetailResponse:
        """
        Drill-down: Lấy chi tiết các hóa đơn còn nợ của một đại lý để kế toán đối chiếu.
        """
        if as_of_date is None:
            as_of_date = datetime.now(TZ_VN).date()

        # Kiểm tra customer tồn tại
        customer = db.query(Customer).filter(Customer.id == customer_id).first()
        if not customer:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Không tìm thấy đại lý với mã ID: {customer_id}"
            )

        # Kiểm tra quyền truy cập (anti-IDOR)
        user_role = getattr(current_user, "role", "")
        if isinstance(user_role, UserRole):
            user_role = user_role.value

        if user_role == UserRole.CUSTOMER.value:
            user_cust_id = getattr(current_user, "customer_id", None)
            if user_cust_id and user_cust_id != customer_id:
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Bạn không có quyền truy cập dữ liệu công nợ của đại lý khác"
                )
            elif not user_cust_id:
                user_cust = db.query(Customer).filter(
                    or_(Customer.email == current_user.email, Customer.phone == current_user.phone_number)
                ).first()
                if not user_cust or user_cust.id != customer_id:
                    raise HTTPException(
                        status_code=status.HTTP_403_FORBIDDEN,
                        detail="Bạn không có quyền truy cập dữ liệu công nợ của đại lý khác"
                    )

        # Lấy thông tin chung của đại lý
        items, _, _ = cls.get_agency_debt_data(
            db=db,
            current_user=current_user,
            as_of_date=as_of_date,
            customer_id=customer_id,
            has_debt_only=False,
        )
        if not items:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Không thể trích xuất dữ liệu công nợ đại lý"
            )
        agency_item = items[0]

        # Lấy danh sách hóa đơn
        invoices = db.query(Invoice).filter(
            Invoice.customer_id == customer_id
        ).order_by(Invoice.issued_at.desc()).all()

        invoice_details: List[AgencyInvoiceDetail] = []
        for inv in invoices:
            aging_info = cls.calculate_invoice_aging(inv, as_of_date)
            inv_detail = AgencyInvoiceDetail(
                invoice_id=inv.id,
                client_request_id=inv.client_request_id,
                issued_at=inv.issued_at,
                due_date=inv.due_date,
                total_amount=int(inv.total_amount or 0),
                paid_amount=int(inv.paid_amount or 0),
                remaining_amount=int(inv.remaining_amount or 0),
                status=inv.status,
                aging_days=aging_info["aging_days"],
                aging_bucket=aging_info["aging_bucket"],
                days_overdue=aging_info["days_overdue"],
                is_overdue=aging_info["is_overdue"],
                notes=inv.notes,
            )
            invoice_details.append(inv_detail)

        return AgencyDebtDetailResponse(
            customer=agency_item,
            invoices=invoice_details,
        )

    @classmethod
    def export_agency_debt_excel(
        cls,
        db: Session,
        current_user: User,
        as_of_date: Optional[date] = None,
        search: Optional[str] = None,
        region: Optional[str] = None,
        customer_group: Optional[str] = None,
        assigned_sales_rep: Optional[str] = None,
        only_overdue: bool = False,
        aging_group: Optional[str] = None,
        has_debt_only: bool = True,
    ) -> io.BytesIO:
        """
        Xuất file Excel báo cáo sổ công nợ đại lý với phân tích tuổi nợ và dòng tổng cộng chân bảng (SCRUM-722).
        Được định dạng chuyên nghiệp chuẩn kế toán với openpyxl.
        """
        items, summary, calculated_date = cls.get_agency_debt_data(
            db=db,
            current_user=current_user,
            as_of_date=as_of_date,
            search=search,
            region=region,
            customer_group=customer_group,
            assigned_sales_rep=assigned_sales_rep,
            only_overdue=only_overdue,
            aging_group=aging_group,
            has_debt_only=has_debt_only,
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "So_Cong_No_Dai_Ly"
        ws.views.sheetView[0].showGridLines = True

        # Định dạng màu sắc và font chữ
        font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
        font_data = Font(name="Calibri", size=11, color="1E293B")
        font_overdue = Font(name="Calibri", size=11, bold=True, color="DC2626")

        fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_header_aging = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
        fill_footer = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        thin_border_side = Side(border_style="thin", color="CBD5E1")
        thick_top_double_bottom = Border(
            top=Side(border_style="thin", color="1E3A8A"),
            bottom=Side(border_style="double", color="1E3A8A"),
            left=thin_border_side,
            right=thin_border_side,
        )
        cell_border = Border(
            top=thin_border_side,
            bottom=thin_border_side,
            left=thin_border_side,
            right=thin_border_side,
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

        num_format_currency = "#,##0"

        # 1. Tiêu đề báo cáo
        ws.merge_cells("A1:P1")
        title_cell = ws["A1"]
        title_cell.value = "BÁO CÁO SỔ CÔNG NỢ & PHÂN TÍCH TUỔI NỢ THEO ĐẠI LÝ"
        title_cell.font = font_title
        title_cell.alignment = align_center
        ws.row_dimensions[1].height = 30

        ws.merge_cells("A2:P2")
        sub_cell = ws["A2"]
        date_str = calculated_date.strftime("%d/%m/%Y")
        now_str = datetime.now(TZ_VN).strftime("%d/%m/%Y %H:%M")
        user_name = getattr(current_user, "full_name", "") or getattr(current_user, "username", "Kế toán")
        sub_cell.value = f"Ngày chốt công nợ: {date_str}  |  Thời điểm lập: {now_str}  |  Người lập: {user_name}"
        sub_cell.font = font_subtitle
        sub_cell.alignment = align_center
        ws.row_dimensions[2].height = 20

        # Ghi chú bộ lọc nếu có
        filter_texts = []
        if search: filter_texts.append(f"Tìm kiếm: '{search}'")
        if region: filter_texts.append(f"Khu vực: {region}")
        if customer_group: filter_texts.append(f"Nhóm: {customer_group}")
        if only_overdue: filter_texts.append("Chỉ lọc đại lý có nợ quá hạn")
        if aging_group: filter_texts.append(f"Nhóm tuổi nợ: {aging_group}")

        ws.merge_cells("A3:P3")
        filter_cell = ws["A3"]
        if filter_texts:
            filter_cell.value = "Bộ lọc đang áp dụng: " + " | ".join(filter_texts)
        else:
            filter_cell.value = "Phạm vi: Toàn bộ đại lý có phát sinh công nợ"
        filter_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
        filter_cell.alignment = align_center
        ws.row_dimensions[3].height = 18

        ws.append([]) # Dòng trống 4

        # 2. Hàng tiêu đề bảng (Row 5)
        headers = [
            ("STT", align_center, 6),
            ("Mã đại lý", align_center, 14),
            ("Tên đại lý", align_left, 30),
            ("Khu vực", align_center, 14),
            ("NVKD phụ trách", align_left, 18),
            ("Hạn mức nợ (VNĐ)", align_right, 18),
            ("Phải thu (VNĐ)", align_right, 18),
            ("Đã thu (VNĐ)", align_right, 18),
            ("Còn nợ (VNĐ)", align_right, 18),
            ("Tuổi nợ 0–30 ngày", align_right, 18),
            ("Tuổi nợ 31–60 ngày", align_right, 18),
            ("Tuổi nợ 61–90 ngày", align_right, 18),
            ("Tuổi nợ >90 ngày", align_right, 18),
            ("Nợ quá hạn (VNĐ)", align_right, 18),
            ("Số ngày trễ max", align_center, 15),
            ("Tình trạng", align_center, 14),
        ]

        row_header_idx = 5
        ws.row_dimensions[row_header_idx].height = 28

        for col_idx, (h_title, _, col_width) in enumerate(headers, start=1):
            cell = ws.cell(row=row_header_idx, column=col_idx, value=h_title)
            cell.font = font_header
            # Nhấn mạnh nhóm cột tuổi nợ bằng màu xanh sáng hơn một chút
            if 10 <= col_idx <= 13:
                cell.fill = fill_header_aging
            else:
                cell.fill = fill_header
            cell.alignment = align_header
            cell.border = cell_border
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width

        # 3. Điền các dòng dữ liệu
        current_row = 6
        for idx, item in enumerate(items, start=1):
            ws.row_dimensions[current_row].height = 22
            row_fill = fill_zebra if idx % 2 == 0 else None

            row_data = [
                (idx, align_center, None, font_data),
                (item.customer_code, align_center, None, font_bold),
                (item.customer_name, align_left, None, font_data),
                (item.region or "-", align_center, None, font_data),
                (item.assigned_sales_rep or "-", align_left, None, font_data),
                (item.credit_limit, align_right, num_format_currency, font_data),
                (item.total_receivable, align_right, num_format_currency, font_data),
                (item.total_paid, align_right, num_format_currency, font_data),
                (item.total_debt, align_right, num_format_currency, font_bold),
                (item.debt_0_30, align_right, num_format_currency, font_data),
                (item.debt_31_60, align_right, num_format_currency, font_data),
                (item.debt_61_90, align_right, num_format_currency, font_data),
                (item.debt_over_90, align_right, num_format_currency, font_overdue if item.debt_over_90 > 0 else font_data),
                (item.overdue_debt, align_right, num_format_currency, font_overdue if item.overdue_debt > 0 else font_data),
                (item.max_overdue_days if item.is_overdue else 0, align_center, None, font_overdue if item.is_overdue else font_data),
                ("Quá hạn" if item.is_overdue else "Trong hạn", align_center, None, font_overdue if item.is_overdue else font_data),
            ]

            for col_idx, (val, alignment, n_format, cell_font) in enumerate(row_data, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.alignment = alignment
                cell.font = cell_font
                cell.border = cell_border
                if n_format:
                    cell.number_format = n_format
                if row_fill:
                    cell.fill = row_fill

            current_row += 1

        # 4. Dòng TỔNG CỘNG ở chân bảng (AC: SCRUM-721 & SCRUM-722)
        ws.row_dimensions[current_row].height = 26
        footer_data = [
            (f"TỔNG CỘNG ({summary.total_customers} đại lý)", align_left, None, font_bold),
            ("", align_center, None, font_bold),
            ("", align_center, None, font_bold),
            ("", align_center, None, font_bold),
            ("", align_center, None, font_bold),
            (0, align_right, num_format_currency, font_bold), # Bỏ trống hạn mức tổng hoặc ghi "-"
            (summary.total_receivable, align_right, num_format_currency, font_bold),
            (summary.total_paid, align_right, num_format_currency, font_bold),
            (summary.total_debt, align_right, num_format_currency, font_bold),
            (summary.total_debt_0_30, align_right, num_format_currency, font_bold),
            (summary.total_debt_31_60, align_right, num_format_currency, font_bold),
            (summary.total_debt_61_90, align_right, num_format_currency, font_bold),
            (summary.total_debt_over_90, align_right, num_format_currency, font_bold),
            (summary.total_overdue_debt, align_right, num_format_currency, font_overdue if summary.total_overdue_debt > 0 else font_bold),
            (f"{summary.overdue_customers_count} ĐL", align_center, None, font_overdue if summary.overdue_customers_count > 0 else font_bold),
            (f"{summary.overdue_customers_count} quá hạn", align_center, None, font_overdue if summary.overdue_customers_count > 0 else font_bold),
        ]

        for col_idx, (val, alignment, n_format, cell_font) in enumerate(footer_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.alignment = alignment
            cell.font = cell_font
            cell.fill = fill_footer
            cell.border = thick_top_double_bottom
            if n_format:
                cell.number_format = n_format

        # Cột F (Hạn mức tổng) hiển thị ký tự "-"
        ws.cell(row=current_row, column=6, value="-")

        # Merge cột A đến E cho chữ TỔNG CỘNG sau khi đã thiết lập style
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)

        # 5. Lưu vào stream
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
