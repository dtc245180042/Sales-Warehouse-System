"""
Service xử lý Sổ chi tiết công nợ đại lý (Story S8-10, SCRUM-844, SCRUM-845, SCRUM-846, SCRUM-847).
Đặc tả nghiệp vụ:
- Tổng hợp toàn bộ phát sinh tăng (Hóa đơn bán hàng) và phát sinh giảm (Phiếu thu tiền, Chứng từ giảm nợ).
- Hỗ trợ lọc theo khoảng thời gian linh hoạt (từ ngày - đến ngày).
- Tự động kết chuyển số dư đầu kỳ (opening balance) chuẩn xác cho các phát sinh trước kỳ.
- Sắp xếp thứ tự thời gian tăng dần và tính số dư luỹ kế (running balance) sau từng giao dịch.
- Xuất tài liệu chuẩn kế toán: Excel (.xlsx), PDF chuẩn vector A4, và HTML in ấn offline.
- Triệt tiêu N+1 queries, tuân thủ nguyên tắc hiệu năng AGENTS.md.
"""

from datetime import datetime, date, time, timezone
import io
import os
import logging
from typing import List, Optional, Tuple, Dict, Any
from zoneinfo import ZoneInfo
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from sqlalchemy import or_

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pymupdf

from app.models.customer import Customer
from app.models.customer_credit_profile import CustomerCreditProfile
from app.models.customer_debt_profile import CustomerDebtProfile
from app.models.invoice import Invoice, InvoiceStatus
from app.models.payment import Payment
from app.models.credit_note import CreditNote
from app.schemas.customer_debt_ledger import (
    DebtLedgerItem,
    CustomerDebtSummary,
    CustomerDebtLedgerResponse,
    CustomerDebtLedgerOverviewItem,
    CustomerDebtLedgerOverviewResponse,
)

logger = logging.getLogger(__name__)
TZ_VN = ZoneInfo("Asia/Ho_Chi_Minh")


def _get_customer_or_404(db: Session, customer_id_or_code: str) -> Customer:
    """Tra cứu đại lý / khách hàng theo ID hoặc Mã khách hàng."""
    customer = db.query(Customer).filter(
        (Customer.id == customer_id_or_code) | (Customer.code == customer_id_or_code)
    ).first()
    if not customer:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Không tìm thấy đại lý / khách hàng với mã '{customer_id_or_code}'."
        )
    return customer


def _parse_date_bounds(
    from_date: Optional[Any],
    to_date: Optional[Any]
) -> Tuple[Optional[datetime], Optional[datetime], Optional[str], Optional[str]]:
    """
    Chuẩn hóa khoảng thời gian lọc sang datetime có timezone UTC để truy vấn DB.
    Trả về: (start_dt, end_dt, from_date_str, to_date_str)
    """
    start_dt: Optional[datetime] = None
    end_dt: Optional[datetime] = None
    from_str: Optional[str] = None
    to_str: Optional[str] = None

    if from_date:
        if isinstance(from_date, str):
            from_str = from_date.strip()
            d = datetime.strptime(from_str, "%Y-%m-%d").date()
        elif isinstance(from_date, (datetime, date)):
            d = from_date if isinstance(from_date, date) else from_date.date()
            from_str = d.strftime("%Y-%m-%d")
        else:
            d = None

        if d:
            # Bắt đầu từ 00:00:00 theo giờ VN
            local_start = datetime.combine(d, time.min, tzinfo=TZ_VN)
            start_dt = local_start.astimezone(timezone.utc)

    if to_date:
        if isinstance(to_date, str):
            to_str = to_date.strip()
            d = datetime.strptime(to_str, "%Y-%m-%d").date()
        elif isinstance(to_date, (datetime, date)):
            d = to_date if isinstance(to_date, date) else to_date.date()
            to_str = d.strftime("%Y-%m-%d")
        else:
            d = None

        if d:
            # Kết thúc lúc 23:59:59.999999 theo giờ VN
            local_end = datetime.combine(d, time.max, tzinfo=TZ_VN)
            end_dt = local_end.astimezone(timezone.utc)

    return start_dt, end_dt, from_str, to_str


def _get_customer_debt_profile_info(db: Session, customer_id: str) -> Tuple[int, int, int]:
    """Lấy hạn mức, số ngày nợ và dư nợ lưu đệm từ profile."""
    credit_prof = db.query(CustomerCreditProfile).filter(
        CustomerCreditProfile.customer_id == customer_id
    ).first()

    credit_limit = 0
    max_debt_days = 0
    current_debt = 0

    if credit_prof:
        credit_limit = int(credit_prof.credit_limit or 0)
        max_debt_days = int(credit_prof.max_debt_days or 0)
        current_debt = int(credit_prof.current_debt or 0)
    else:
        debt_prof = db.query(CustomerDebtProfile).filter(
            CustomerDebtProfile.customer_id == customer_id
        ).first()
        if debt_prof:
            credit_limit = int(round(debt_prof.credit_limit or 0))
            current_debt = int(round(debt_prof.current_debt or 0))

    return credit_limit, max_debt_days, current_debt


def get_customer_debt_ledger(
    db: Session,
    customer_id_or_code: str,
    from_date: Optional[Any] = None,
    to_date: Optional[Any] = None,
) -> CustomerDebtLedgerResponse:
    """
    Xây dựng logic tổng hợp sổ chi tiết công nợ theo thứ tự thời gian và tính số dư luỹ kế (SCRUM-844, SCRUM-845).
    """
    customer = _get_customer_or_404(db, customer_id_or_code)
    start_dt, end_dt, from_str, to_str = _parse_date_bounds(from_date, to_date)
    credit_limit, max_debt_days, cached_debt = _get_customer_debt_profile_info(db, customer.id)

    # 1. Truy vấn các chứng từ kế toán của đại lý
    invoices = db.query(Invoice).filter(Invoice.customer_id == customer.id).all()
    payments = db.query(Payment).filter(Payment.customer_id == customer.id).all()
    credit_notes = db.query(CreditNote).filter(CreditNote.customer_id == customer.id).all()

    # 2. Chuẩn hóa tất cả chứng từ thành cấu trúc thô để phân tích
    raw_txns = []

    # A. Hóa đơn bán hàng (Tăng nợ)
    for inv in invoices:
        txn_dt = inv.issued_at or inv.created_at
        if txn_dt and txn_dt.tzinfo is None:
            txn_dt = txn_dt.replace(tzinfo=timezone.utc)

        ref = None
        if inv.items and len(inv.items) > 0:
            ref = inv.items[0].order_id or inv.items[0].dispatch_id

        raw_txns.append({
            "id": inv.id,
            "date": txn_dt,
            "document_type": "INVOICE",
            "document_type_label": "Hóa đơn bán hàng",
            "document_code": inv.id,
            "reference_code": ref,
            "description": inv.notes or f"Hóa đơn xuất bán hàng theo phiếu xuất",
            "increase_amount": int(inv.total_amount or 0),
            "decrease_amount": 0,
            "due_date": inv.due_date,
            "status": inv.status or InvoiceStatus.UNPAID.value,
            "priority": 1,  # Hóa đơn phát sinh trước nếu trùng timestamp
        })

    # B. Phiếu thu tiền (Giảm nợ)
    for pay in payments:
        txn_dt = pay.payment_date or pay.created_at
        if txn_dt and txn_dt.tzinfo is None:
            txn_dt = txn_dt.replace(tzinfo=timezone.utc)

        method_vn = "Chuyển khoản" if pay.payment_method == "BANK_TRANSFER" else ("Tiền mặt" if pay.payment_method == "CASH" else "Thẻ")
        raw_txns.append({
            "id": pay.id,
            "date": txn_dt,
            "document_type": "PAYMENT",
            "document_type_label": "Phiếu thu tiền",
            "document_code": pay.id,
            "reference_code": pay.reference_code,
            "description": pay.notes or f"Thu tiền bán hàng ({method_vn})",
            "increase_amount": 0,
            "decrease_amount": int(pay.amount or 0),
            "due_date": None,
            "status": "COMPLETED",
            "priority": 3,
        })

    # C. Chứng từ giảm trừ công nợ khi trả hàng (Giảm nợ)
    for cn in credit_notes:
        txn_dt = cn.adjustment_date or cn.created_at
        if txn_dt and txn_dt.tzinfo is None:
            txn_dt = txn_dt.replace(tzinfo=timezone.utc)

        ref = cn.return_order_id or cn.invoice_id
        raw_txns.append({
            "id": cn.id,
            "date": txn_dt,
            "document_type": "CREDIT_NOTE",
            "document_type_label": "Chứng từ giảm nợ",
            "document_code": cn.id,
            "reference_code": ref,
            "description": cn.notes or f"Giảm trừ công nợ theo phiếu trả hàng {cn.return_order_id}",
            "increase_amount": 0,
            "decrease_amount": int(cn.adjustment_amount or 0),
            "due_date": None,
            "status": "ADJUSTED",
            "priority": 2,
        })

    # 3. Phân tách giao dịch theo khoảng thời gian và tính số dư đầu kỳ
    opening_balance = 0
    in_period_txns = []

    for t in raw_txns:
        t_date = t["date"]
        inc = t["increase_amount"]
        dec = t["decrease_amount"]

        if start_dt and t_date < start_dt:
            # Thuộc về kỳ trước -> tích lũy vào số dư đầu kỳ
            opening_balance += (inc - dec)
        elif end_dt and t_date > end_dt:
            # Nằm sau kỳ lọc -> bỏ qua
            continue
        else:
            # Nằm trong kỳ đối chiếu
            in_period_txns.append(t)

    # 4. Sắp xếp thứ tự thời gian tăng dần: date ASC, priority ASC, id ASC
    in_period_txns.sort(key=lambda x: (x["date"], x["priority"], x["id"]))

    # 5. Tính số dư luỹ kế (running balance) qua từng dòng
    current_running_balance = opening_balance
    ledger_items: List[DebtLedgerItem] = []

    total_increase = 0
    total_decrease = 0

    for item in in_period_txns:
        inc = item["increase_amount"]
        dec = item["decrease_amount"]
        total_increase += inc
        total_decrease += dec
        current_running_balance += (inc - dec)

        ledger_items.append(
            DebtLedgerItem(
                id=item["id"],
                date=item["date"],
                document_type=item["document_type"],
                document_type_label=item["document_type_label"],
                document_code=item["document_code"],
                reference_code=item["reference_code"],
                description=item["description"],
                increase_amount=inc,
                decrease_amount=dec,
                running_balance=current_running_balance,
                due_date=item["due_date"],
                status=item["status"],
            )
        )

    # 6. Số dư cuối kỳ = Số dư đầu kỳ + Phát sinh Tăng - Phát sinh Giảm
    closing_balance = opening_balance + total_increase - total_decrease

    # Tính dư nợ toàn bộ hiện tại của khách hàng (không phụ thuộc filter)
    all_increase = sum(t["increase_amount"] for t in raw_txns)
    all_decrease = sum(t["decrease_amount"] for t in raw_txns)
    lifetime_debt = all_increase - all_decrease
    available_credit = max(0, credit_limit - lifetime_debt)

    customer_summary = CustomerDebtSummary(
        customer_id=customer.id,
        customer_code=customer.code or customer.id,
        customer_name=customer.name or customer.code or "Đại lý",
        phone=customer.phone,
        email=customer.email,
        address=customer.address,
        tax_code=getattr(customer, "tax_code", None),
        credit_limit=credit_limit,
        max_debt_days=max_debt_days,
        current_debt=lifetime_debt,
        available_credit=available_credit,
    )

    return CustomerDebtLedgerResponse(
        customer=customer_summary,
        from_date=from_str,
        to_date=to_str,
        opening_balance=opening_balance,
        total_increase=total_increase,
        total_decrease=total_decrease,
        closing_balance=closing_balance,
        items=ledger_items,
        total_items=len(ledger_items),
    )


def export_customer_debt_ledger_excel(
    db: Session,
    customer_id_or_code: str,
    from_date: Optional[Any] = None,
    to_date: Optional[Any] = None,
) -> bytes:
    """
    Phát triển chức năng xuất Excel sổ chi tiết công nợ của đại lý (SCRUM-847).
    Tạo bảng tính Excel (.xlsx) chuẩn biểu mẫu kế toán Việt Nam (Mẫu S31-DN - TT 200/2014/TT-BTC).
    """
    ledger = get_customer_debt_ledger(db, customer_id_or_code, from_date, to_date)
    cust = ledger.customer

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SoChiTietCongNo"
    ws.views.sheetView[0].showGridLines = True

    # Định nghĩa Font và Màu sắc chuẩn doanh nghiệp
    font_company = Font(name="Segoe UI", size=9, bold=True, color="334155")
    font_sub_header = Font(name="Segoe UI", size=9, italic=True, color="64748b")
    font_title = Font(name="Segoe UI", size=15, bold=True, color="1e3a8a")
    font_period = Font(name="Segoe UI", size=10, italic=True, color="475569")
    font_info_label = Font(name="Segoe UI", size=9.5, bold=True, color="1e293b")
    font_info_val = Font(name="Segoe UI", size=9.5, color="0f172a")
    font_tbl_header = Font(name="Segoe UI", size=10, bold=True, color="ffffff")
    font_data = Font(name="Segoe UI", size=9.5, color="0f172a")
    font_bold = Font(name="Segoe UI", size=9.5, bold=True, color="0f172a")
    font_closing = Font(name="Segoe UI", size=10, bold=True, color="1e3a8a")

    fill_header = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
    fill_opening = PatternFill(start_color="f8fafc", end_color="f8fafc", fill_type="solid")
    fill_even = PatternFill(start_color="f8fafc", end_color="f8fafc", fill_type="solid")
    fill_summary = PatternFill(start_color="f1f5f9", end_color="f1f5f9", fill_type="solid")
    fill_closing = PatternFill(start_color="e0e7ff", end_color="e0e7ff", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="cbd5e1")
    thick_border_side = Side(border_style="medium", color="1e40af")
    double_border_side = Side(border_style="double", color="334155")

    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_summary = Border(left=thin_border_side, right=thin_border_side, top=Side(border_style="thin", color="475569"), bottom=double_border_side)

    # 1. Header cơ quan / doanh nghiệp
    ws["A1"] = "CÔNG TY CỔ PHẦN PHÂN PHỐI BÁN HÀNG & KHO HÀNG"
    ws["A1"].font = font_company
    ws["A2"] = "Địa chỉ: Số 123 Đường Công Nghệ, Khu Công Nghệ Cao, TP. Hà Nội"
    ws["A2"].font = font_sub_header

    ws["I1"] = "Mẫu số S31-DN"
    ws["I1"].font = Font(name="Segoe UI", size=9, bold=True, color="334155")
    ws["I1"].alignment = Alignment(horizontal="right")
    ws["I2"] = "(Ban hành theo TT số 200/2014/TT-BTC)"
    ws["I2"].font = font_sub_header
    ws["I2"].alignment = Alignment(horizontal="right")

    # 2. Tiêu đề Sổ chi tiết
    ws.merge_cells("A4:I4")
    ws["A4"] = "SỔ CHI TIẾT CÔNG NỢ PHẢI THU (TÀI KHOẢN 131)"
    ws["A4"].font = font_title
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 28

    period_text = "Toàn bộ thời gian phát sinh"
    if ledger.from_date and ledger.to_date:
        period_text = f"Từ ngày {ledger.from_date} đến ngày {ledger.to_date}"
    elif ledger.from_date:
        period_text = f"Từ ngày {ledger.from_date} đến nay"
    elif ledger.to_date:
        period_text = f"Đến ngày {ledger.to_date}"

    ws.merge_cells("A5:I5")
    ws["A5"] = f"Kỳ đối chiếu: {period_text}"
    ws["A5"].font = font_period
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

    # 3. Thông tin đại lý
    ws["A7"] = "Tên đại lý / Đối tác:"
    ws["A7"].font = font_info_label
    ws["B7"] = cust.customer_name
    ws["B7"].font = font_bold

    ws["F7"] = "Mã đại lý:"
    ws["F7"].font = font_info_label
    ws["G7"] = cust.customer_code or cust.customer_id
    ws["G7"].font = font_bold

    ws["A8"] = "Địa chỉ đăng ký:"
    ws["A8"].font = font_info_label
    ws["B8"] = cust.address or "Chưa cập nhật"
    ws["B8"].font = font_info_val

    ws["F8"] = "Số điện thoại:"
    ws["F8"].font = font_info_label
    ws["G8"] = cust.phone or "Chưa cập nhật"
    ws["G8"].font = font_info_val

    ws["A9"] = "Mã số thuế:"
    ws["A9"].font = font_info_label
    ws["B9"] = cust.tax_code or "Chưa có"
    ws["B9"].font = font_info_val

    ws["F9"] = "Hạn mức công nợ:"
    ws["F9"].font = font_info_label
    ws["G9"] = f"{cust.credit_limit:,.0f} đ (Thời hạn: {cust.max_debt_days} ngày)" if cust.credit_limit > 0 else "Không cấp hạn mức"
    ws["G9"].font = font_bold

    # 4. Tiêu đề bảng Sổ công nợ
    start_row = 11
    headers = [
        ("STT", 6, "center"),
        ("Ngày chứng từ", 15, "center"),
        ("Số chứng từ", 16, "center"),
        ("Loại chứng từ", 18, "center"),
        ("Diễn giải giao dịch", 32, "left"),
        ("Mã tham chiếu", 16, "center"),
        ("Phát sinh Tăng (Nợ)", 18, "right"),
        ("Phát sinh Giảm (Có)", 18, "right"),
        ("Số dư nợ luỹ kế", 20, "right"),
    ]

    ws.row_dimensions[start_row].height = 25
    for col_idx, (h_name, width, align) in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 5. Dòng số dư đầu kỳ
    current_r = start_row + 1
    ws.row_dimensions[current_r].height = 20
    ws.cell(row=current_r, column=1, value="").border = border_cell
    ws.cell(row=current_r, column=2, value="").border = border_cell
    ws.cell(row=current_r, column=3, value="").border = border_cell
    ws.cell(row=current_r, column=4, value="").border = border_cell

    cell_open_desc = ws.cell(row=current_r, column=5, value="SỐ DƯ ĐẦU KỲ")
    cell_open_desc.font = Font(name="Segoe UI", size=9.5, bold=True, italic=True, color="1e40af")
    cell_open_desc.border = border_cell
    cell_open_desc.fill = fill_opening

    ws.cell(row=current_r, column=6, value="").border = border_cell
    ws.cell(row=current_r, column=7, value="").border = border_cell
    ws.cell(row=current_r, column=8, value="").border = border_cell

    cell_open_val = ws.cell(row=current_r, column=9, value=ledger.opening_balance)
    cell_open_val.font = font_bold
    cell_open_val.number_format = "#,##0"
    cell_open_val.alignment = Alignment(horizontal="right", vertical="center")
    cell_open_val.border = border_cell
    cell_open_val.fill = fill_opening

    # 6. Các dòng chứng từ phát sinh trong kỳ
    for idx, item in enumerate(ledger.items, 1):
        current_r += 1
        ws.row_dimensions[current_r].height = 19
        row_fill = fill_even if idx % 2 == 0 else PatternFill(fill_type=None)

        # Định dạng ngày theo DD/MM/YYYY HH:mm
        dt_vn = item.date.astimezone(TZ_VN) if item.date.tzinfo else item.date
        date_str = dt_vn.strftime("%d/%m/%Y %H:%M")

        row_data = [
            (idx, "center", None),
            (date_str, "center", None),
            (item.document_code, "center", font_bold),
            (item.document_type_label, "left", None),
            (item.description or "", "left", None),
            (item.reference_code or "", "center", None),
            (item.increase_amount if item.increase_amount > 0 else None, "right", None),
            (item.decrease_amount if item.decrease_amount > 0 else None, "right", None),
            (item.running_balance, "right", font_bold),
        ]

        for col_idx, (val, align, custom_font) in enumerate(row_data, 1):
            cell = ws.cell(row=current_r, column=col_idx, value=val)
            cell.font = custom_font or font_data
            cell.alignment = Alignment(horizontal=align, vertical="center")
            cell.border = border_cell
            if row_fill.fill_type:
                cell.fill = row_fill

            if col_idx in [7, 8, 9] and val is not None:
                cell.number_format = "#,##0"

    # 7. Dòng Tổng cộng phát sinh trong kỳ
    current_r += 1
    ws.row_dimensions[current_r].height = 22
    for c in range(1, 7):
        ws.cell(row=current_r, column=c, value="").border = border_summary
        ws.cell(row=current_r, column=c).fill = fill_summary

    cell_sum_label = ws.cell(row=current_r, column=5, value="TỔNG PHÁT SINH TRONG KỲ")
    cell_sum_label.font = font_bold
    cell_sum_label.alignment = Alignment(horizontal="right", vertical="center")

    cell_inc = ws.cell(row=current_r, column=7, value=ledger.total_increase)
    cell_inc.font = font_bold
    cell_inc.number_format = "#,##0"
    cell_inc.alignment = Alignment(horizontal="right", vertical="center")
    cell_inc.border = border_summary
    cell_inc.fill = fill_summary

    cell_dec = ws.cell(row=current_r, column=8, value=ledger.total_decrease)
    cell_dec.font = font_bold
    cell_dec.number_format = "#,##0"
    cell_dec.alignment = Alignment(horizontal="right", vertical="center")
    cell_dec.border = border_summary
    cell_dec.fill = fill_summary

    ws.cell(row=current_r, column=9, value="").border = border_summary
    ws.cell(row=current_r, column=9).fill = fill_summary

    # 8. Dòng Số dư nợ cuối kỳ
    current_r += 1
    ws.row_dimensions[current_r].height = 24
    for c in range(1, 9):
        ws.cell(row=current_r, column=c, value="").border = border_cell
        ws.cell(row=current_r, column=c).fill = fill_closing

    cell_close_label = ws.cell(row=current_r, column=5, value="SỐ DƯ NỢ CUỐI KỲ PHẢI THU")
    cell_close_label.font = font_closing
    cell_close_label.alignment = Alignment(horizontal="right", vertical="center")

    cell_close = ws.cell(row=current_r, column=9, value=ledger.closing_balance)
    cell_close.font = font_closing
    cell_close.number_format = "#,##0"
    cell_close.alignment = Alignment(horizontal="right", vertical="center")
    cell_close.border = border_cell
    cell_close.fill = fill_closing

    # 9. Phần chữ ký xác nhận đối chiếu công nợ
    current_r += 2
    ws.cell(row=current_r, column=2, value="NGƯỜI LẬP BIỂU").font = font_bold
    ws.cell(row=current_r, column=2).alignment = Alignment(horizontal="center")

    ws.cell(row=current_r, column=4, value="KẾ TOÁN CÔNG NỢ").font = font_bold
    ws.cell(row=current_r, column=4).alignment = Alignment(horizontal="center")

    ws.cell(row=current_r, column=7, value="KẾ TOÁN TRƯỞNG").font = font_bold
    ws.cell(row=current_r, column=7).alignment = Alignment(horizontal="center")

    ws.cell(row=current_r, column=9, value="ĐẠI DIỆN ĐẠI LÝ").font = font_bold
    ws.cell(row=current_r, column=9).alignment = Alignment(horizontal="center")

    current_r += 1
    ws.cell(row=current_r, column=2, value="(Ký, họ tên)").font = font_sub_header
    ws.cell(row=current_r, column=2).alignment = Alignment(horizontal="center")

    ws.cell(row=current_r, column=4, value="(Ký, họ tên)").font = font_sub_header
    ws.cell(row=current_r, column=4).alignment = Alignment(horizontal="center")

    ws.cell(row=current_r, column=7, value="(Ký, họ tên)").font = font_sub_header
    ws.cell(row=current_r, column=7).alignment = Alignment(horizontal="center")

    ws.cell(row=current_r, column=9, value="(Ký, đóng dấu, họ tên)").font = font_sub_header
    ws.cell(row=current_r, column=9).alignment = Alignment(horizontal="center")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def export_customer_debt_ledger_pdf(
    db: Session,
    customer_id_or_code: str,
    from_date: Optional[Any] = None,
    to_date: Optional[Any] = None,
) -> bytes:
    """
    Phát triển chức năng xuất PDF sổ chi tiết công nợ của đại lý (SCRUM-846).
    Tạo tài liệu PDF khổ A4 Landscape chuẩn vector với font tiếng Việt sắc nét.
    """
    ledger = get_customer_debt_ledger(db, customer_id_or_code, from_date, to_date)
    cust = ledger.customer

    # A4 Landscape: 842 x 595 points
    page_width = 842.0
    page_height = 595.0
    margin_x = 36.0
    margin_y = 36.0

    doc = pymupdf.open()
    page = doc.new_page(width=page_width, height=page_height)

    # Nạp font tiếng Việt hệ thống nếu có
    font_name = "helv"
    font_bold_name = "helv-bold"
    font_path_candidates = [
        os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "arial.ttf"),
        os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "times.ttf"),
        os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "segoeui.ttf"),
    ]
    font_file = None
    for p in font_path_candidates:
        if os.path.exists(p):
            font_file = p
            break

    if font_file:
        try:
            page.insert_font(fontname="VNFont", fontfile=font_file)
            font_name = "VNFont"
            font_bold_name = "VNFont"
        except Exception as e:
            logger.warning(f"Could not load system font for PDF: {e}")

    # Màu sắc định dạng
    color_navy = (0.12, 0.25, 0.69)       # #1e40af
    color_dark = (0.06, 0.09, 0.16)       # #0f172a
    color_gray = (0.39, 0.45, 0.55)       # #64748b
    color_border = (0.80, 0.84, 0.88)     # #cbd5e1
    color_bg_header = (0.12, 0.25, 0.69)
    color_bg_even = (0.97, 0.98, 0.99)
    color_bg_highlight = (0.94, 0.96, 1.0)

    # 1. Header công ty & Mẫu biểu
    page.insert_text(
        (margin_x, margin_y + 10),
        "CÔNG TY CỔ PHẦN PHÂN PHỐI BÁN HÀNG & KHO HÀNG",
        fontname=font_bold_name,
        fontsize=9,
        color=color_navy,
    )
    page.insert_text(
        (margin_x, margin_y + 22),
        "Địa chỉ: Khu Công Nghệ Cao, TP. Hà Nội • Hotline: 1900 6868",
        fontname=font_name,
        fontsize=8,
        color=color_gray,
    )

    page.insert_text(
        (page_width - margin_x - 130, margin_y + 10),
        "Mẫu số S31-DN",
        fontname=font_bold_name,
        fontsize=9,
        color=color_dark,
    )
    page.insert_text(
        (page_width - margin_x - 170, margin_y + 22),
        "(Ban hành theo TT số 200/2014/TT-BTC)",
        fontname=font_name,
        fontsize=8,
        color=color_gray,
    )

    # 2. Tiêu đề chính
    title_text = "SỔ CHI TIẾT CÔNG NỢ ĐẠI LÝ (TK 131)"
    page.insert_text(
        (page_width / 2 - 140, margin_y + 45),
        title_text,
        fontname=font_bold_name,
        fontsize=14,
        color=color_navy,
    )

    period_str = "Toàn bộ thời gian phát sinh"
    if ledger.from_date and ledger.to_date:
        period_str = f"Kỳ đối chiếu: Từ ngày {ledger.from_date} đến ngày {ledger.to_date}"
    elif ledger.from_date:
        period_str = f"Kỳ đối chiếu: Từ ngày {ledger.from_date} đến nay"
    elif ledger.to_date:
        period_str = f"Kỳ đối chiếu: Đến ngày {ledger.to_date}"

    page.insert_text(
        (page_width / 2 - 110, margin_y + 60),
        period_str,
        fontname=font_name,
        fontsize=9,
        color=color_gray,
    )

    # 3. Hộp thông tin đối tác & Hạn mức
    info_y = margin_y + 70
    info_rect = pymupdf.Rect(margin_x, info_y, page_width - margin_x, info_y + 36)
    page.draw_rect(info_rect, color=color_border, fill=color_bg_even, width=0.8)

    # Cột 1
    page.insert_text((margin_x + 10, info_y + 14), f"Tên đại lý: {cust.customer_name}", fontname=font_bold_name, fontsize=8.5, color=color_dark)
    page.insert_text((margin_x + 10, info_y + 27), f"Địa chỉ: {cust.address or 'Chưa cập nhật'}", fontname=font_name, fontsize=8, color=color_gray)

    # Cột 2
    col2_x = margin_x + 360
    page.insert_text((col2_x, info_y + 14), f"Mã đại lý: {cust.customer_code or cust.customer_id}", fontname=font_bold_name, fontsize=8.5, color=color_dark)
    page.insert_text((col2_x, info_y + 27), f"Điện thoại: {cust.phone or 'Chưa cập nhật'}", fontname=font_name, fontsize=8, color=color_gray)

    # Cột 3
    col3_x = margin_x + 550
    limit_str = f"{cust.credit_limit:,.0f} đ" if cust.credit_limit > 0 else "0 đ"
    page.insert_text((col3_x, info_y + 14), f"Hạn mức nợ: {limit_str}", fontname=font_bold_name, fontsize=8.5, color=color_navy)
    page.insert_text((col3_x, info_y + 27), f"Thời hạn nợ: {cust.max_debt_days} ngày", fontname=font_name, fontsize=8, color=color_gray)

    # 4. Bảng chi tiết các cột
    # Cấu hình chiều rộng các cột (Tổng cộng = page_width - 2 * margin_x = 770)
    col_widths = [28, 80, 80, 85, 175, 75, 75, 75, 97]
    col_x = [margin_x]
    for w in col_widths:
        col_x.append(col_x[-1] + w)

    table_top = info_y + 44
    row_height = 18

    # Vẽ Header bảng
    header_rect = pymupdf.Rect(margin_x, table_top, page_width - margin_x, table_top + row_height)
    page.draw_rect(header_rect, color=color_navy, fill=color_bg_header, width=0.8)

    col_headers = ["STT", "Ngày CT", "Số CT", "Loại chứng từ", "Diễn giải giao dịch", "Tham chiếu", "Tăng (Nợ)", "Giảm (Có)", "Dư nợ luỹ kế"]
    for i, h_text in enumerate(col_headers):
        text_x = col_x[i] + 4
        page.insert_text((text_x, table_top + 12), h_text, fontname=font_bold_name, fontsize=8, color=(1, 1, 1))

    # Dòng Số dư đầu kỳ
    current_y = table_top + row_height
    open_rect = pymupdf.Rect(margin_x, current_y, page_width - margin_x, current_y + row_height)
    page.draw_rect(open_rect, color=color_border, fill=color_bg_even, width=0.5)

    page.insert_text((col_x[4] + 4, current_y + 12), "SỐ DƯ ĐẦU KỲ", fontname=font_bold_name, fontsize=8, color=color_navy)
    open_bal_str = f"{ledger.opening_balance:,.0f} đ"
    page.insert_text((col_x[8] + 10, current_y + 12), open_bal_str, fontname=font_bold_name, fontsize=8, color=color_dark)

    current_y += row_height

    # In các dòng chứng từ
    max_rows_per_page = 16
    for idx, item in enumerate(ledger.items, 1):
        if current_y + row_height > page_height - margin_y - 50:
            # Sang trang mới nếu vượt quá giới hạn
            page = doc.new_page(width=page_width, height=page_height)
            if font_file:
                page.insert_font(fontname="VNFont", fontfile=font_file)
            current_y = margin_y + 10

        r_rect = pymupdf.Rect(margin_x, current_y, page_width - margin_x, current_y + row_height)
        bg_color = color_bg_even if idx % 2 == 0 else (1, 1, 1)
        page.draw_rect(r_rect, color=color_border, fill=bg_color, width=0.5)

        dt_vn = item.date.astimezone(TZ_VN) if item.date.tzinfo else item.date
        dt_str = dt_vn.strftime("%d/%m/%y %H:%M")

        inc_str = f"{item.increase_amount:,.0f}" if item.increase_amount > 0 else "-"
        dec_str = f"{item.decrease_amount:,.0f}" if item.decrease_amount > 0 else "-"
        run_str = f"{item.running_balance:,.0f} đ"

        page.insert_text((col_x[0] + 6, current_y + 12), str(idx), fontname=font_name, fontsize=7.5, color=color_dark)
        page.insert_text((col_x[1] + 4, current_y + 12), dt_str, fontname=font_name, fontsize=7.5, color=color_dark)
        page.insert_text((col_x[2] + 4, current_y + 12), item.document_code, fontname=font_bold_name, fontsize=7.5, color=color_navy)
        page.insert_text((col_x[3] + 4, current_y + 12), item.document_type_label[:18], fontname=font_name, fontsize=7.5, color=color_dark)

        # Cắt ngắn diễn giải nếu quá dài
        desc_short = (item.description or "")[:35]
        page.insert_text((col_x[4] + 4, current_y + 12), desc_short, fontname=font_name, fontsize=7.5, color=color_dark)
        page.insert_text((col_x[5] + 4, current_y + 12), (item.reference_code or "-")[:12], fontname=font_name, fontsize=7.5, color=color_gray)

        page.insert_text((col_x[6] + 4, current_y + 12), inc_str, fontname=font_name, fontsize=7.5, color=color_dark)
        page.insert_text((col_x[7] + 4, current_y + 12), dec_str, fontname=font_name, fontsize=7.5, color=color_dark)
        page.insert_text((col_x[8] + 4, current_y + 12), run_str, fontname=font_bold_name, fontsize=7.5, color=color_navy)

        current_y += row_height

    # Dòng Tổng phát sinh trong kỳ
    sum_rect = pymupdf.Rect(margin_x, current_y, page_width - margin_x, current_y + row_height)
    page.draw_rect(sum_rect, color=color_border, fill=color_bg_highlight, width=0.8)

    page.insert_text((col_x[4] + 4, current_y + 12), "TỔNG PHÁT SINH TRONG KỲ", fontname=font_bold_name, fontsize=8, color=color_dark)
    page.insert_text((col_x[6] + 4, current_y + 12), f"{ledger.total_increase:,.0f} đ", fontname=font_bold_name, fontsize=8, color=color_dark)
    page.insert_text((col_x[7] + 4, current_y + 12), f"{ledger.total_decrease:,.0f} đ", fontname=font_bold_name, fontsize=8, color=color_dark)

    current_y += row_height

    # Dòng Số dư cuối kỳ
    close_rect = pymupdf.Rect(margin_x, current_y, page_width - margin_x, current_y + row_height)
    page.draw_rect(close_rect, color=color_navy, fill=color_bg_highlight, width=1.0)

    page.insert_text((col_x[4] + 4, current_y + 12), "SỐ DƯ NỢ CUỐI KỲ PHẢI THU", fontname=font_bold_name, fontsize=8.5, color=color_navy)
    page.insert_text((col_x[8] + 4, current_y + 12), f"{ledger.closing_balance:,.0f} đ", fontname=font_bold_name, fontsize=8.5, color=color_navy)

    current_y += row_height + 15

    # 5. Khu vực chữ ký đối chiếu 2 bên
    if current_y + 40 <= page_height - margin_y:
        sig_y = current_y
        page.insert_text((margin_x + 30, sig_y), "NGƯỜI LẬP BIỂU", fontname=font_bold_name, fontsize=8, color=color_dark)
        page.insert_text((margin_x + 35, sig_y + 12), "(Ký, họ tên)", fontname=font_name, fontsize=7.5, color=color_gray)

        page.insert_text((margin_x + 230, sig_y), "KẾ TOÁN CÔNG NỢ", fontname=font_bold_name, fontsize=8, color=color_dark)
        page.insert_text((margin_x + 238, sig_y + 12), "(Ký, họ tên)", fontname=font_name, fontsize=7.5, color=color_gray)

        page.insert_text((margin_x + 440, sig_y), "KẾ TOÁN TRƯỞNG", fontname=font_bold_name, fontsize=8, color=color_dark)
        page.insert_text((margin_x + 448, sig_y + 12), "(Ký, họ tên)", fontname=font_name, fontsize=7.5, color=color_gray)

        page.insert_text((margin_x + 630, sig_y), "ĐẠI DIỆN ĐẠI LÝ", fontname=font_bold_name, fontsize=8, color=color_dark)
        page.insert_text((margin_x + 625, sig_y + 12), "(Ký, đóng dấu)", fontname=font_name, fontsize=7.5, color=color_gray)

    return doc.tobytes()


def generate_customer_debt_ledger_html(
    db: Session,
    customer_id_or_code: str,
    from_date: Optional[Any] = None,
    to_date: Optional[Any] = None,
) -> str:
    """
    Sinh tài liệu HTML hoàn chỉnh chuẩn in A4 Landscape Offline-First.
    Sẵn sàng mở trong Tab mới hoặc IFrame để Kế toán bấm Ctrl+P / In PDF ngay tức thì.
    """
    ledger = get_customer_debt_ledger(db, customer_id_or_code, from_date, to_date)
    cust = ledger.customer

    period_str = "Toàn bộ thời gian phát sinh"
    if ledger.from_date and ledger.to_date:
        period_str = f"Từ ngày {ledger.from_date} đến ngày {ledger.to_date}"
    elif ledger.from_date:
        period_str = f"Từ ngày {ledger.from_date} đến nay"
    elif ledger.to_date:
        period_str = f"Đến ngày {ledger.to_date}"

    rows_html = ""
    for idx, item in enumerate(ledger.items, 1):
        dt_vn = item.date.astimezone(TZ_VN) if item.date.tzinfo else item.date
        dt_str = dt_vn.strftime("%d/%m/%Y %H:%M")

        inc_td = f"{item.increase_amount:,.0f} đ" if item.increase_amount > 0 else "-"
        dec_td = f"{item.decrease_amount:,.0f} đ" if item.decrease_amount > 0 else "-"
        run_td = f"{item.running_balance:,.0f} đ"

        badge_class = "badge-blue" if item.document_type == "INVOICE" else ("badge-green" if item.document_type == "PAYMENT" else "badge-amber")

        rows_html += f"""
        <tr>
          <td class="text-center">{idx}</td>
          <td class="text-center">{dt_str}</td>
          <td class="text-center font-bold text-navy">{item.document_code}</td>
          <td><span class="badge {badge_class}">{item.document_type_label}</span></td>
          <td>{item.description or ''}</td>
          <td class="text-center text-muted">{item.reference_code or '-'}</td>
          <td class="text-right text-rose font-medium">{inc_td}</td>
          <td class="text-right text-emerald font-medium">{dec_td}</td>
          <td class="text-right font-bold text-navy">{run_td}</td>
        </tr>
        """

    html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
  <meta charset="UTF-8">
  <title>Sổ chi tiết công nợ - {cust.customer_name}</title>
  <style>
    @page {{
      size: A4 landscape;
      margin: 10mm 12mm;
    }}
    body {{
      font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
      color: #0f172a;
      background: #ffffff;
      margin: 0;
      padding: 20px;
      font-size: 11px;
    }}
    .no-print {{
      margin-bottom: 20px;
      display: flex;
      justify-content: flex-end;
      gap: 10px;
    }}
    .btn {{
      padding: 8px 16px;
      border-radius: 6px;
      font-weight: 600;
      cursor: pointer;
      border: 1px solid transparent;
      font-size: 12px;
    }}
    .btn-primary {{
      background-color: #1e40af;
      color: white;
    }}
    .btn-secondary {{
      background-color: #f1f5f9;
      color: #334155;
      border-color: #cbd5e1;
    }}
    @media print {{
      .no-print {{ display: none !important; }}
      body {{ padding: 0; }}
    }}
    .header-table {{
      width: 100%;
      border-collapse: collapse;
      margin-bottom: 12px;
    }}
    .header-table td {{
      vertical-align: top;
    }}
    .title-box {{
      text-align: center;
      margin: 10px 0 16px 0;
    }}
    .main-title {{
      font-size: 16px;
      font-weight: 800;
      color: #1e3a8a;
      text-transform: uppercase;
      letter-spacing: 0.5px;
    }}
    .sub-title {{
      font-style: italic;
      color: #64748b;
      margin-top: 3px;
    }}
    .info-card {{
      background-color: #f8fafc;
      border: 1px solid #e2e8f0;
      border-radius: 8px;
      padding: 10px 14px;
      margin-bottom: 14px;
      display: grid;
      grid-template-columns: 2fr 1fr 1fr;
      gap: 12px;
    }}
    .info-group p {{
      margin: 2px 0;
    }}
    .data-table {{
      width: 100%;
      border-collapse: collapse;
      font-size: 10.5px;
    }}
    .data-table th {{
      background-color: #1e40af;
      color: white;
      font-weight: 700;
      padding: 8px 6px;
      border: 1px solid #1e3a8a;
      text-align: center;
    }}
    .data-table td {{
      padding: 6px 8px;
      border: 1px solid #cbd5e1;
      vertical-align: middle;
    }}
    .data-table tr:nth-child(even) {{
      background-color: #f8fafc;
    }}
    .opening-row {{
      background-color: #f1f5f9 !important;
      font-weight: bold;
    }}
    .summary-row {{
      background-color: #e0e7ff !important;
      font-weight: bold;
      border-top: 2px solid #334155;
    }}
    .closing-row {{
      background-color: #dbeafe !important;
      font-weight: 800;
      font-size: 11px;
      color: #1e3a8a;
    }}
    .text-center {{ text-align: center; }}
    .text-right {{ text-align: right; }}
    .font-bold {{ font-weight: bold; }}
    .text-navy {{ color: #1e40af; }}
    .text-rose {{ color: #e11d48; }}
    .text-emerald {{ color: #059669; }}
    .text-muted {{ color: #64748b; }}
    .badge {{
      display: inline-block;
      padding: 2px 6px;
      border-radius: 4px;
      font-size: 9px;
      font-weight: 600;
    }}
    .badge-blue {{ background-color: #dbeafe; color: #1e40af; }}
    .badge-green {{ background-color: #d1fae5; color: #065f46; }}
    .badge-amber {{ background-color: #fef3c7; color: #92400e; }}
    .signature-grid {{
      margin-top: 30px;
      display: grid;
      grid-template-columns: repeat(4, 1fr);
      text-align: center;
      page-break-inside: avoid;
    }}
    .sig-title {{
      font-weight: 700;
      color: #0f172a;
    }}
    .sig-subtitle {{
      font-style: italic;
      color: #64748b;
      font-size: 9.5px;
      margin-top: 3px;
    }}
  </style>
</head>
<body>
  <div class="no-print">
    <button class="btn btn-secondary" onclick="window.close()">Đóng</button>
    <button class="btn btn-primary" onclick="window.print()">In sổ chi tiết / Lưu PDF</button>
  </div>

  <table class="header-table">
    <tr>
      <td>
        <strong style="color: #1e3a8a; font-size: 11px;">CÔNG TY CỔ PHẦN PHÂN PHỐI BÁN HÀNG & KHO HÀNG</strong><br/>
        <span style="color: #64748b;">Khu Công Nghệ Cao, TP. Hà Nội • Hotline: 1900 6868</span>
      </td>
      <td style="text-align: right;">
        <strong>Mẫu số S31-DN</strong><br/>
        <span style="color: #64748b; font-style: italic;">(Ban hành theo TT số 200/2014/TT-BTC)</span>
      </td>
    </tr>
  </table>

  <div class="title-box">
    <div class="main-title">SỔ CHI TIẾT CÔNG NỢ PHẢI THU (TÀI KHOẢN 131)</div>
    <div class="sub-title">Kỳ đối chiếu: {period_str}</div>
  </div>

  <div class="info-card">
    <div class="info-group">
      <p><strong>Tên đại lý:</strong> {cust.customer_name}</p>
      <p><strong>Địa chỉ:</strong> {cust.address or 'Chưa cập nhật'}</p>
      <p><strong>Mã số thuế:</strong> {cust.tax_code or 'Chưa có'}</p>
    </div>
    <div class="info-group">
      <p><strong>Mã đại lý:</strong> {cust.customer_code or cust.customer_id}</p>
      <p><strong>Điện thoại:</strong> {cust.phone or 'Chưa cập nhật'}</p>
      <p><strong>Email:</strong> {cust.email or 'Chưa cập nhật'}</p>
    </div>
    <div class="info-group">
      <p><strong>Hạn mức nợ:</strong> {cust.credit_limit:,.0f} đ</p>
      <p><strong>Thời hạn nợ:</strong> {cust.max_debt_days} ngày</p>
      <p><strong>Dư nợ hiện tại:</strong> <span class="text-rose font-bold">{cust.current_debt:,.0f} đ</span></p>
    </div>
  </div>

  <table class="data-table">
    <thead>
      <tr>
        <th style="width: 35px;">STT</th>
        <th style="width: 100px;">Ngày chứng từ</th>
        <th style="width: 100px;">Số chứng từ</th>
        <th style="width: 110px;">Loại chứng từ</th>
        <th>Diễn giải nghiệp vụ</th>
        <th style="width: 90px;">Mã tham chiếu</th>
        <th style="width: 105px;">Phát sinh Tăng (Nợ)</th>
        <th style="width: 105px;">Phát sinh Giảm (Có)</th>
        <th style="width: 115px;">Số dư nợ luỹ kế</th>
      </tr>
    </thead>
    <tbody>
      <tr class="opening-row">
        <td></td>
        <td></td>
        <td></td>
        <td></td>
        <td class="font-bold text-navy" style="font-style: italic;">SỐ DƯ ĐẦU KỲ</td>
        <td></td>
        <td></td>
        <td></td>
        <td class="text-right font-bold text-navy">{ledger.opening_balance:,.0f} đ</td>
      </tr>
      {rows_html}
      <tr class="summary-row">
        <td colspan="4"></td>
        <td class="text-right font-bold">TỔNG CỘNG PHÁT SINH TRONG KỲ:</td>
        <td></td>
        <td class="text-right text-rose font-bold">{ledger.total_increase:,.0f} đ</td>
        <td class="text-right text-emerald font-bold">{ledger.total_decrease:,.0f} đ</td>
        <td></td>
      </tr>
      <tr class="closing-row">
        <td colspan="4"></td>
        <td class="text-right font-bold">SỐ DƯ NỢ CUỐI KỲ PHẢI THU:</td>
        <td colspan="3"></td>
        <td class="text-right font-bold text-navy" style="font-size: 12px;">{ledger.closing_balance:,.0f} đ</td>
      </tr>
    </tbody>
  </table>

  <div class="signature-grid">
    <div>
      <div class="sig-title">NGƯỜI LẬP BIỂU</div>
      <div class="sig-subtitle">(Ký, ghi rõ họ tên)</div>
    </div>
    <div>
      <div class="sig-title">KẾ TOÁN CÔNG NỢ</div>
      <div class="sig-subtitle">(Ký, ghi rõ họ tên)</div>
    </div>
    <div>
      <div class="sig-title">KẾ TOÁN TRƯỞNG</div>
      <div class="sig-subtitle">(Ký, ghi rõ họ tên)</div>
    </div>
    <div>
      <div class="sig-title">ĐẠI DIỆN ĐẠI LÝ</div>
      <div class="sig-subtitle">(Ký, đóng dấu, ghi rõ họ tên)</div>
    </div>
  </div>
</body>
</html>"""
    return html


def get_all_customers_debt_ledger_overview(
    db: Session,
    search: Optional[str] = None,
    from_date: Optional[Any] = None,
    to_date: Optional[Any] = None,
    page: int = 1,
    page_size: int = 50,
) -> CustomerDebtLedgerOverviewResponse:
    """
    Lấy danh sách tổng quan công nợ các đại lý (hỗ trợ kế toán xem nhanh tổng hợp).
    Triệt tiêu N+1 Query theo nguyên tắc AGENTS.md.
    """
    start_dt, end_dt, from_str, to_str = _parse_date_bounds(from_date, to_date)

    query = db.query(Customer)
    if search:
        search_kw = f"%{search.strip()}%"
        query = query.filter(
            or_(
                Customer.name.ilike(search_kw),
                Customer.code.ilike(search_kw),
                Customer.phone.ilike(search_kw),
            )
        )

    total_customers = query.count()
    customers = query.offset((page - 1) * page_size).limit(page_size).all()
    customer_ids = [c.id for c in customers]

    if not customer_ids:
        return CustomerDebtLedgerOverviewResponse(items=[], total=0, total_closing_debt=0)

    # 1. Gom tất cả profile theo danh sách ID
    profiles = db.query(CustomerCreditProfile).filter(
        CustomerCreditProfile.customer_id.in_(customer_ids)
    ).all()
    profile_map = {p.customer_id: p for p in profiles}

    # 2. Gom tất cả hóa đơn của các đại lý này
    invoices = db.query(Invoice).filter(Invoice.customer_id.in_(customer_ids)).all()
    payments = db.query(Payment).filter(Payment.customer_id.in_(customer_ids)).all()
    credit_notes = db.query(CreditNote).filter(CreditNote.customer_id.in_(customer_ids)).all()

    # Gom theo customer_id
    inv_map: Dict[str, List[Invoice]] = {}
    for inv in invoices:
        inv_map.setdefault(inv.customer_id, []).append(inv)

    pay_map: Dict[str, List[Payment]] = {}
    for pay in payments:
        pay_map.setdefault(pay.customer_id, []).append(pay)

    cn_map: Dict[str, List[CreditNote]] = {}
    for cn in credit_notes:
        cn_map.setdefault(cn.customer_id, []).append(cn)

    items: List[CustomerDebtLedgerOverviewItem] = []
    total_all_closing = 0

    for c in customers:
        c_invs = inv_map.get(c.id, [])
        c_pays = pay_map.get(c.id, [])
        c_cns = cn_map.get(c.id, [])

        p = profile_map.get(c.id)
        limit = int(p.credit_limit or 0) if p else 0
        days = int(p.max_debt_days or 0) if p else 0

        opening = 0
        inc_period = 0
        dec_period = 0
        txn_count = 0

        # Quét hóa đơn
        for inv in c_invs:
            dt = inv.issued_at or inv.created_at
            if dt and dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            amt = int(inv.total_amount or 0)

            if start_dt and dt < start_dt:
                opening += amt
            elif end_dt and dt > end_dt:
                continue
            else:
                inc_period += amt
                txn_count += 1

        # Quét phiếu thu
        for pay in c_pays:
            dt = pay.payment_date or pay.created_at
            if dt and dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            amt = int(pay.amount or 0)

            if start_dt and dt < start_dt:
                opening -= amt
            elif end_dt and dt > end_dt:
                continue
            else:
                dec_period += amt
                txn_count += 1

        # Quét credit note
        for cn in c_cns:
            dt = cn.adjustment_date or cn.created_at
            if dt and dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            amt = int(cn.adjustment_amount or 0)

            if start_dt and dt < start_dt:
                opening -= amt
            elif end_dt and dt > end_dt:
                continue
            else:
                dec_period += amt
                txn_count += 1

        closing = opening + inc_period - dec_period
        total_all_closing += closing

        items.append(
            CustomerDebtLedgerOverviewItem(
                customer_id=c.id,
                customer_code=c.code or c.id,
                customer_name=c.name or c.code or "Đại lý",
                phone=c.phone,
                credit_limit=limit,
                max_debt_days=days,
                opening_balance=opening,
                total_increase=inc_period,
                total_decrease=dec_period,
                closing_balance=closing,
                transactions_count=txn_count,
            )
        )

    return CustomerDebtLedgerOverviewResponse(
        items=items,
        total=total_customers,
        total_closing_debt=total_all_closing,
    )
