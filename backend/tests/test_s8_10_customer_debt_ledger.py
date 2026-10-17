"""
Bộ kiểm thử tự động toàn diện Sổ chi tiết công nợ đại lý (Story S8-10, SCRUM-843 -> SCRUM-849).
Kiểm tra:
- SCRUM-844: API xuất sổ chi tiết công nợ, hỗ trợ lọc theo khoảng thời gian
- SCRUM-845: Logic tổng hợp sổ chi tiết theo thứ tự thời gian và tính số dư luỹ kế
- SCRUM-846: Chức năng xuất PDF sổ chi tiết công nợ chuẩn vector A4
- SCRUM-847: Chức năng xuất Excel sổ chi tiết công nợ chuẩn TT 200/2014/TT-BTC
- SCRUM-849: Kiểm thử đối chiếu số dư đầu kỳ, phát sinh tăng giảm và số dư cuối kỳ
"""

from datetime import datetime, timedelta, timezone
import io
import pytest
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, lay_phien_db
from app.core.security import bam_mat_khau, tao_token_truy_cap
from app.models.auth import User, UserRole
from app.models.customer import Customer
from app.models.customer_credit_profile import CustomerCreditProfile
from app.models.invoice import Invoice, InvoiceStatus
from app.models.payment import Payment, PaymentMethod
from app.models.credit_note import CreditNote
from main import app

# Database SQLite in-memory tách biệt dùng StaticPool
TEST_DB_URL = "sqlite:///:memory:"
test_engine = create_engine(
    TEST_DB_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_lay_phien_db():
    phien_db = TestingSessionLocal()
    try:
        yield phien_db
    finally:
        phien_db.close()


@pytest.fixture(autouse=True)
def setup_ledger_test_database():
    """Tạo schema và dữ liệu người dùng trước mỗi test."""
    Base.metadata.create_all(bind=test_engine)
    app.dependency_overrides[lay_phien_db] = override_lay_phien_db

    db = TestingSessionLocal()
    users = [
        User(
            id=1,
            username="admin_user",
            email="admin@test.local",
            role=UserRole.ADMIN.value,
            hashed_password=bam_mat_khau("Pass@1234"),
            is_active=True,
            token_version=1,
        ),
        User(
            id=2,
            username="accountant_user",
            email="accountant@test.local",
            role=UserRole.ACCOUNTANT.value,
            hashed_password=bam_mat_khau("Pass@1234"),
            is_active=True,
            token_version=1,
        ),
        User(
            id=3,
            username="customer_user_1",
            email="cust1@test.local",
            role=UserRole.CUSTOMER.value,
            customer_id="CUST-LEDGER-01",
            hashed_password=bam_mat_khau("Pass@1234"),
            is_active=True,
            token_version=1,
        ),
        User(
            id=4,
            username="customer_user_2",
            email="cust2@test.local",
            role=UserRole.CUSTOMER.value,
            customer_id="CUST-LEDGER-02",
            hashed_password=bam_mat_khau("Pass@1234"),
            is_active=True,
            token_version=1,
        ),
    ]
    for u in users:
        db.add(u)

    # Thêm khách hàng mẫu
    cust1 = Customer(
        id="CUST-LEDGER-01",
        code="DL-HA-NOI",
        name="Đại Lý Phân Phối Hà Nội",
        phone="0912345678",
        email="hanoi@agency.vn",
        address="100 Phố Huế, Hai Bà Trưng, Hà Nội",
        tax_code="0101234567",
        status="active",
    )
    cust2 = Customer(
        id="CUST-LEDGER-02",
        code="DL-DA-NANG",
        name="Đại Lý Thiết Bị Số Đà Nẵng",
        phone="0987654321",
        email="danang@agency.vn",
        address="200 Nguyễn Văn Linh, Hải Châu, Đà Nẵng",
        tax_code="0409876543",
        status="active",
    )
    db.add(cust1)
    db.add(cust2)

    # Hồ sơ tín dụng
    cp1 = CustomerCreditProfile(
        customer_id="CUST-LEDGER-01",
        credit_limit=100000000,
        max_debt_days=30,
        current_debt=0,
    )
    db.add(cp1)
    db.commit()
    db.close()

    yield

    Base.metadata.drop_all(bind=test_engine)
    app.dependency_overrides.clear()


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def auth_header_accountant():
    token = tao_token_truy_cap({"sub": "accountant_user", "role": "Accountant", "user_id": 2, "token_version": 1})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def auth_header_admin():
    token = tao_token_truy_cap({"sub": "admin_user", "role": "Admin", "user_id": 1, "token_version": 1})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def auth_header_customer1():
    token = tao_token_truy_cap({"sub": "customer_user_1", "role": "Customer", "user_id": 3, "customer_id": "CUST-LEDGER-01", "token_version": 1})
    return {"Authorization": f"Bearer {token}"}


def test_empty_ledger(client, auth_header_accountant):
    """Đại lý chưa có giao dịch nào trả về số dư 0 và danh sách rỗng."""
    resp = client.get(
        "/api/v1/customer-debt-ledgers/CUST-LEDGER-01",
        headers=auth_header_accountant,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["customer"]["customer_id"] == "CUST-LEDGER-01"
    assert data["customer"]["customer_name"] == "Đại Lý Phân Phối Hà Nội"
    assert data["opening_balance"] == 0
    assert data["total_increase"] == 0
    assert data["total_decrease"] == 0
    assert data["closing_balance"] == 0
    assert len(data["items"]) == 0


def test_running_balance_calculation_scrum_845(client, auth_header_accountant):
    """
    SCRUM-845: Kiểm tra tính toán số dư luỹ kế (running balance) qua từng dòng:
    - T0: Hóa đơn 1 (+50M) -> Dư nợ: 50M
    - T1: Hóa đơn 2 (+30M) -> Dư nợ: 80M
    - T2: Phiếu thu 1 (-40M) -> Dư nợ: 40M
    - T3: Giảm nợ 1 (-10M) -> Dư nợ: 30M
    - Cuối kỳ: 30M
    """
    db = TestingSessionLocal()
    now = datetime.now(timezone.utc)

    # 1. Hóa đơn 1 (+50,000,000)
    inv1 = Invoice(
        id="HD-2026-0000001",
        customer_id="CUST-LEDGER-01",
        customer_name="Đại Lý Phân Phối Hà Nội",
        subtotal_amount=50000000,
        total_amount=50000000,
        paid_amount=0,
        remaining_amount=50000000,
        status=InvoiceStatus.UNPAID.value,
        issued_at=now - timedelta(days=10),
        due_date=now + timedelta(days=20),
        issued_by=1,
    )
    # 2. Hóa đơn 2 (+30,000,000)
    inv2 = Invoice(
        id="HD-2026-0000002",
        customer_id="CUST-LEDGER-01",
        customer_name="Đại Lý Phân Phối Hà Nội",
        subtotal_amount=30000000,
        total_amount=30000000,
        paid_amount=0,
        remaining_amount=30000000,
        status=InvoiceStatus.UNPAID.value,
        issued_at=now - timedelta(days=7),
        due_date=now + timedelta(days=23),
        issued_by=1,
    )
    # 3. Phiếu thu 1 (-40,000,000)
    pay1 = Payment(
        id="PT-2026-0000001",
        customer_id="CUST-LEDGER-01",
        customer_name="Đại Lý Phân Phối Hà Nội",
        amount=40000000,
        payment_method=PaymentMethod.BANK_TRANSFER.value,
        payment_date=now - timedelta(days=5),
        collector_id=2,
        reference_code="VCB-123456",
        notes="Thu tiền đợt 1",
    )
    # 4. Chứng từ giảm nợ 1 (-10,000,000)
    cn1 = CreditNote(
        id="CN-2026-0000001",
        return_order_id="RO-2026-00001",
        invoice_id="HD-2026-0000001",
        customer_id="CUST-LEDGER-01",
        customer_name="Đại Lý Phân Phối Hà Nội",
        adjustment_amount=10000000,
        adjustment_date=now - timedelta(days=2),
        created_by=2,
        notes="Giảm trừ hàng trả về",
    )
    db.add_all([inv1, inv2, pay1, cn1])
    db.commit()
    db.close()

    resp = client.get(
        "/api/v1/customer-debt-ledgers/CUST-LEDGER-01",
        headers=auth_header_accountant,
    )
    assert resp.status_code == 200
    data = resp.json()

    assert data["opening_balance"] == 0
    assert data["total_increase"] == 80000000
    assert data["total_decrease"] == 50000000
    assert data["closing_balance"] == 30000000
    assert len(data["items"]) == 4

    # Kiểm tra số dư luỹ kế từng dòng
    items = data["items"]
    # Dòng 1: HD1
    assert items[0]["document_code"] == "HD-2026-0000001"
    assert items[0]["increase_amount"] == 50000000
    assert items[0]["decrease_amount"] == 0
    assert items[0]["running_balance"] == 50000000

    # Dòng 2: HD2
    assert items[1]["document_code"] == "HD-2026-0000002"
    assert items[1]["increase_amount"] == 30000000
    assert items[1]["decrease_amount"] == 0
    assert items[1]["running_balance"] == 80000000

    # Dòng 3: PT1
    assert items[2]["document_code"] == "PT-2026-0000001"
    assert items[2]["increase_amount"] == 0
    assert items[2]["decrease_amount"] == 40000000
    assert items[2]["running_balance"] == 40000000

    # Dòng 4: CN1
    assert items[3]["document_code"] == "CN-2026-0000001"
    assert items[3]["increase_amount"] == 0
    assert items[3]["decrease_amount"] == 10000000
    assert items[3]["running_balance"] == 30000000


def test_date_range_filtering_and_opening_balance_scrum_844(client, auth_header_accountant):
    """
    SCRUM-844: Kiểm tra lọc theo khoảng thời gian và tính số dư đầu kỳ:
    - Ngày 01/10: Hóa đơn 1 (+50M)
    - Ngày 05/10: Phiếu thu 1 (-20M)
      --> Trước ngày 10/10: Số dư đầu kỳ = 50M - 20M = 30M
    - Ngày 15/10: Hóa đơn 2 (+40M)
    - Ngày 20/10: Phiếu thu 2 (-15M)
      --> Trong kỳ (10/10 -> 25/10): Tăng 40M, Giảm 15M
      --> Cuối kỳ = 30M + 40M - 15M = 55M
    - Ngày 28/10 (sau kỳ): Hóa đơn 3 (+10M) -> Bị loại trừ khỏi kỳ lọc!
    """
    db = TestingSessionLocal()
    d1 = datetime(2026, 10, 1, 10, 0, tzinfo=timezone.utc)
    d5 = datetime(2026, 10, 5, 14, 0, tzinfo=timezone.utc)
    d15 = datetime(2026, 10, 15, 9, 30, tzinfo=timezone.utc)
    d20 = datetime(2026, 10, 20, 16, 0, tzinfo=timezone.utc)
    d28 = datetime(2026, 10, 28, 11, 0, tzinfo=timezone.utc)

    inv1 = Invoice(
        id="HD-2026-1001",
        customer_id="CUST-LEDGER-01",
        total_amount=50000000,
        subtotal_amount=50000000,
        issued_at=d1,
        due_date=d1 + timedelta(days=30),
        issued_by=1,
    )
    pay1 = Payment(
        id="PT-2026-1005",
        customer_id="CUST-LEDGER-01",
        amount=20000000,
        payment_date=d5,
        collector_id=2,
    )
    inv2 = Invoice(
        id="HD-2026-1015",
        customer_id="CUST-LEDGER-01",
        total_amount=40000000,
        subtotal_amount=40000000,
        issued_at=d15,
        due_date=d15 + timedelta(days=30),
        issued_by=1,
    )
    pay2 = Payment(
        id="PT-2026-1020",
        customer_id="CUST-LEDGER-01",
        amount=15000000,
        payment_date=d20,
        collector_id=2,
    )
    inv3 = Invoice(
        id="HD-2026-1028",
        customer_id="CUST-LEDGER-01",
        total_amount=10000000,
        subtotal_amount=10000000,
        issued_at=d28,
        due_date=d28 + timedelta(days=30),
        issued_by=1,
    )
    db.add_all([inv1, pay1, inv2, pay2, inv3])
    db.commit()
    db.close()

    # Lọc từ ngày 2026-10-10 đến 2026-10-25
    resp = client.get(
        "/api/v1/customer-debt-ledgers/CUST-LEDGER-01?from_date=2026-10-10&to_date=2026-10-25",
        headers=auth_header_accountant,
    )
    assert resp.status_code == 200
    data = resp.json()

    assert data["from_date"] == "2026-10-10"
    assert data["to_date"] == "2026-10-25"
    assert data["opening_balance"] == 30000000
    assert data["total_increase"] == 40000000
    assert data["total_decrease"] == 15000000
    assert data["closing_balance"] == 55000000

    # Chỉ có 2 chứng từ trong kỳ (HD 15/10 và PT 20/10)
    items = data["items"]
    assert len(items) == 2
    assert items[0]["document_code"] == "HD-2026-1015"
    assert items[0]["running_balance"] == 70000000  # 30M đầu kỳ + 40M = 70M
    assert items[1]["document_code"] == "PT-2026-1020"
    assert items[1]["running_balance"] == 55000000  # 70M - 15M = 55M


def test_export_excel_scrum_847(client, auth_header_accountant):
    """SCRUM-847: Kiểm tra xuất file Excel chuẩn .xlsx có đầy đủ dữ liệu và sheet cấu trúc."""
    db = TestingSessionLocal()
    now = datetime.now(timezone.utc)
    inv = Invoice(
        id="HD-2026-EXCEL01",
        customer_id="CUST-LEDGER-01",
        total_amount=25000000,
        subtotal_amount=25000000,
        issued_at=now,
        due_date=now + timedelta(days=30),
        issued_by=1,
    )
    db.add(inv)
    db.commit()
    db.close()

    resp = client.get(
        "/api/v1/customer-debt-ledgers/CUST-LEDGER-01/export-excel",
        headers=auth_header_accountant,
    )
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
    assert "attachment; filename=" in resp.headers["content-disposition"]

    # Đọc lại bằng openpyxl để kiểm tra tính hợp lệ của file Excel
    content = resp.content
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "SoChiTietCongNo" in wb.sheetnames
    ws = wb["SoChiTietCongNo"]
    assert "SỔ CHI TIẾT CÔNG NỢ PHẢI THU" in str(ws["A4"].value)
    assert "Đại Lý Phân Phối Hà Nội" in str(ws["B7"].value)


def test_export_pdf_scrum_846(client, auth_header_accountant):
    """SCRUM-846: Kiểm tra xuất file PDF chuẩn có magic header %PDF và nội dung."""
    db = TestingSessionLocal()
    now = datetime.now(timezone.utc)
    inv = Invoice(
        id="HD-2026-PDF01",
        customer_id="CUST-LEDGER-01",
        total_amount=15000000,
        subtotal_amount=15000000,
        issued_at=now,
        due_date=now + timedelta(days=30),
        issued_by=1,
    )
    db.add(inv)
    db.commit()
    db.close()

    resp = client.get(
        "/api/v1/customer-debt-ledgers/CUST-LEDGER-01/export-pdf",
        headers=auth_header_accountant,
    )
    assert resp.status_code == 200
    assert "application/pdf" in resp.headers["content-type"]
    assert "attachment; filename=" in resp.headers["content-disposition"]
    assert resp.content.startswith(b"%PDF-")


def test_print_html(client, auth_header_accountant):
    """Kiểm tra xuất trang HTML in ấn hoàn chỉnh offline-first."""
    resp = client.get(
        "/api/v1/customer-debt-ledgers/CUST-LEDGER-01/print-html",
        headers=auth_header_accountant,
    )
    assert resp.status_code == 200
    assert "text/html" in resp.headers["content-type"]
    assert "SỔ CHI TIẾT CÔNG NỢ PHẢI THU" in resp.text
    assert "Đại Lý Phân Phối Hà Nội" in resp.text


def test_overview_endpoint(client, auth_header_accountant):
    """Kiểm tra API danh sách tổng quan công nợ các đại lý."""
    resp = client.get(
        "/api/v1/customer-debt-ledgers?search=Hà+Nội",
        headers=auth_header_accountant,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["total"] >= 1
    assert any(i["customer_id"] == "CUST-LEDGER-01" for i in data["items"])


def test_customer_not_found(client, auth_header_accountant):
    """Truy cập mã đại lý không tồn tại trả về 404 Not Found."""
    resp = client.get(
        "/api/v1/customer-debt-ledgers/NON-EXISTENT-ID",
        headers=auth_header_accountant,
    )
    assert resp.status_code == 404
    assert "Không tìm thấy" in resp.json()["detail"]


def test_customer_cannot_access_backoffice_ledger(client, auth_header_customer1):
    """
    Theo cơ chế Auto-Discovery và bảo mật phân quyền Default-Deny của AGENTS.md,
    tất cả các router backoffice nội bộ đều tự động chặn vai trò Customer (trả về 403 Forbidden).
    """
    resp = client.get(
        "/api/v1/customer-debt-ledgers/CUST-LEDGER-01",
        headers=auth_header_customer1,
    )
    assert resp.status_code == 403
