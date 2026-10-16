import io
from datetime import datetime, date, timedelta, timezone
from zoneinfo import ZoneInfo
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import openpyxl

from app.core.database import Base, lay_phien_db
from app.core.security import bam_mat_khau, tao_token_truy_cap
from app.models.auth import User, UserRole
from app.models.customer import Customer
from app.models.customer_credit_profile import CustomerCreditProfile
from app.models.invoice import Invoice, InvoiceStatus
from app.services.agency_debt_service import AgencyDebtService
from main import app

TEST_DB_URL = "sqlite:///:memory:"
test_engine = create_engine(
    TEST_DB_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_lay_phien_db():
    phien_db = TestingSessionLocal()
    try:
        yield phien_db
    finally:
        phien_db.close()


def auth_header(user_id: int, username: str, role: str) -> dict:
    token = tao_token_truy_cap(du_lieu={
        "sub": str(user_id),
        "user_id": user_id,
        "username": username,
        "role": role,
        "token_version": 1,
    })
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(autouse=True)
def setup_test_db():
    """Tạo schema và dữ liệu mẫu trước mỗi test."""
    Base.metadata.create_all(bind=test_engine)
    app.dependency_overrides[lay_phien_db] = override_lay_phien_db

    db = TestingSessionLocal()
    # 1. Tạo users
    users = [
        User(
            id=1,
            username="admin_debt",
            email="admin_debt@test.local",
            role=UserRole.ADMIN.value,
            hashed_password=bam_mat_khau("Pass@1234"),
            is_active=True,
            token_version=1,
        ),
        User(
            id=2,
            username="accountant_debt",
            email="accountant_debt@test.local",
            role=UserRole.ACCOUNTANT.value,
            hashed_password=bam_mat_khau("Pass@1234"),
            is_active=True,
            token_version=1,
        ),
        User(
            id=3,
            username="sales_rep_debt",
            email="sales_rep@test.local",
            role=UserRole.SALES_REP.value,
            hashed_password=bam_mat_khau("Pass@1234"),
            is_active=True,
            token_version=1,
        ),
        User(
            id=4,
            username="customer_user",
            email="cust_a@test.local",
            phone_number="0901111111",
            role=UserRole.CUSTOMER.value,
            customer_id="CUS-A",
            hashed_password=bam_mat_khau("Pass@1234"),
            is_active=True,
            token_version=1,
        ),
    ]
    for u in users:
        db.add(u)

    # 2. Tạo 2 đại lý
    cust_a = Customer(
        id="CUS-A",
        code="DL-001",
        name="Đại Lý Tuấn Phương",
        phone="0901111111",
        email="cust_a@test.local",
        region="Miền Bắc",
        customer_group="TIER_1",
        assigned_sales_rep="sales_rep_debt",
    )
    cust_b = Customer(
        id="CUS-B",
        code="DL-002",
        name="Đại Lý Hải Đăng",
        phone="0902222222",
        email="cust_b@test.local",
        region="Miền Nam",
        customer_group="WHOLESALE",
        assigned_sales_rep="sales_mgr",
    )
    db.add_all([cust_a, cust_b])

    # 3. Hạn mức công nợ
    cp_a = CustomerCreditProfile(customer_id="CUS-A", credit_limit=50000000, max_debt_days=30)
    cp_b = CustomerCreditProfile(customer_id="CUS-B", credit_limit=100000000, max_debt_days=45)
    db.add_all([cp_a, cp_b])

    # 4. Hóa đơn cho Đại Lý A:
    # Inv 1: 10M, trả 2M -> còn 8M (tuổi 10 ngày -> 0-30 ngày, trong hạn)
    # Inv 2: 5M, trả 0 -> còn 5M (tuổi 40 ngày -> 31-60 ngày, quá hạn 10 ngày)
    inv_a1 = Invoice(
        id="HD-A1",
        customer_id="CUS-A",
        issued_at=datetime(2026, 10, 5, tzinfo=timezone.utc),
        due_date=datetime(2026, 10, 25, tzinfo=timezone.utc),
        total_amount=10000000,
        paid_amount=2000000,
        remaining_amount=8000000,
        status=InvoiceStatus.PARTIALLY_PAID.value,
        issued_by=1,
    )
    inv_a2 = Invoice(
        id="HD-A2",
        customer_id="CUS-A",
        issued_at=datetime(2026, 9, 5, tzinfo=timezone.utc),
        due_date=datetime(2026, 10, 5, tzinfo=timezone.utc),
        total_amount=5000000,
        paid_amount=0,
        remaining_amount=5000000,
        status=InvoiceStatus.UNPAID.value,
        issued_by=1,
    )

    # 5. Hóa đơn cho Đại Lý B:
    # Inv 3: 20M, trả 5M -> còn 15M (tuổi 70 ngày -> 61-90 ngày, quá hạn)
    # Inv 4: 30M, trả 0 -> còn 30M (tuổi 105 ngày -> >90 ngày, quá hạn)
    inv_b1 = Invoice(
        id="HD-B1",
        customer_id="CUS-B",
        issued_at=datetime(2026, 8, 6, tzinfo=timezone.utc),
        due_date=datetime(2026, 9, 6, tzinfo=timezone.utc),
        total_amount=20000000,
        paid_amount=5000000,
        remaining_amount=15000000,
        status=InvoiceStatus.PARTIALLY_PAID.value,
        issued_by=1,
    )
    inv_b2 = Invoice(
        id="HD-B2",
        customer_id="CUS-B",
        issued_at=datetime(2026, 7, 2, tzinfo=timezone.utc),
        due_date=datetime(2026, 8, 2, tzinfo=timezone.utc),
        total_amount=30000000,
        paid_amount=0,
        remaining_amount=30000000,
        status=InvoiceStatus.UNPAID.value,
        issued_by=1,
    )
    db.add_all([inv_a1, inv_a2, inv_b1, inv_b2])
    db.commit()
    db.close()

    yield

    Base.metadata.drop_all(bind=test_engine)


def test_01_calculate_invoice_aging_logic_scrum717():
    """
    SCRUM-717: Kiểm thử logic phân chia 4 nhóm tuổi nợ (0-30, 31-60, 61-90, >90)
    và xác định nợ quá hạn theo ngày đến hạn.
    """
    as_of = date(2026, 10, 15)

    # 1. Hóa đơn 10 ngày tuổi, còn trong hạn
    inv1 = Invoice(
        id="HD-TEST-001",
        customer_id="CUS-001",
        issued_at=datetime(2026, 10, 5, tzinfo=timezone.utc),
        due_date=datetime(2026, 10, 25, tzinfo=timezone.utc),
        total_amount=10000000,
        paid_amount=0,
        remaining_amount=10000000,
    )
    res1 = AgencyDebtService.calculate_invoice_aging(inv1, as_of)
    assert res1["aging_days"] == 10
    assert res1["aging_bucket"] == "0_30"
    assert res1["is_overdue"] is False
    assert res1["days_overdue"] == 0
    assert res1["overdue_amount"] == 0

    # 2. Hóa đơn 45 ngày tuổi, quá hạn 15 ngày
    inv2 = Invoice(
        id="HD-TEST-002",
        customer_id="CUS-001",
        issued_at=datetime(2026, 8, 31, tzinfo=timezone.utc),
        due_date=datetime(2026, 9, 30, tzinfo=timezone.utc),
        total_amount=20000000,
        paid_amount=5000000,
        remaining_amount=15000000,
    )
    res2 = AgencyDebtService.calculate_invoice_aging(inv2, as_of)
    assert res2["aging_days"] == 45
    assert res2["aging_bucket"] == "31_60"
    assert res2["is_overdue"] is True
    assert res2["days_overdue"] == 15
    assert res2["overdue_amount"] == 15000000

    # 3. Hóa đơn 75 ngày tuổi (nhóm 61-90)
    inv3 = Invoice(
        id="HD-TEST-003",
        customer_id="CUS-001",
        issued_at=datetime(2026, 8, 1, tzinfo=timezone.utc),
        due_date=datetime(2026, 8, 31, tzinfo=timezone.utc),
        total_amount=30000000,
        paid_amount=0,
        remaining_amount=30000000,
    )
    res3 = AgencyDebtService.calculate_invoice_aging(inv3, as_of)
    assert res3["aging_days"] == 75
    assert res3["aging_bucket"] == "61_90"
    assert res3["is_overdue"] is True

    # 4. Hóa đơn 120 ngày tuổi (nhóm over_90)
    inv4 = Invoice(
        id="HD-TEST-004",
        customer_id="CUS-001",
        issued_at=datetime(2026, 6, 17, tzinfo=timezone.utc),
        due_date=datetime(2026, 7, 17, tzinfo=timezone.utc),
        total_amount=40000000,
        paid_amount=0,
        remaining_amount=40000000,
    )
    res4 = AgencyDebtService.calculate_invoice_aging(inv4, as_of)
    assert res4["aging_days"] == 120
    assert res4["aging_bucket"] == "over_90"
    assert res4["is_overdue"] is True

    # 5. Hóa đơn quá hạn nhưng đã thanh toán hết (remaining_amount = 0) -> không tính quá hạn
    inv5 = Invoice(
        id="HD-TEST-005",
        customer_id="CUS-001",
        issued_at=datetime(2026, 6, 1, tzinfo=timezone.utc),
        due_date=datetime(2026, 7, 1, tzinfo=timezone.utc),
        total_amount=50000000,
        paid_amount=50000000,
        remaining_amount=0,
    )
    res5 = AgencyDebtService.calculate_invoice_aging(inv5, as_of)
    assert res5["is_overdue"] is False
    assert res5["overdue_amount"] == 0


def test_02_agency_debt_report_api_and_summary_scrum718_721():
    """
    SCRUM-718 & SCRUM-721: API báo cáo sổ công nợ theo đại lý với phân tích tuổi nợ
    và tính toán chính xác dòng tổng cộng ở chân bảng.
    """
    client = TestClient(app)
    headers = auth_header(2, "accountant_debt", UserRole.ACCOUNTANT.value)

    res = client.get(
        "/api/v1/agency-debt-reports?as_of_date=2026-10-15",
        headers=headers,
    )
    assert res.status_code == 200
    data = res.json()

    assert data["total"] == 2
    items = data["items"]
    item_a = next(x for x in items if x["customer_code"] == "DL-001")
    item_b = next(x for x in items if x["customer_code"] == "DL-002")

    # Kiểm tra số liệu Đại lý A
    assert item_a["total_receivable"] == 15000000
    assert item_a["total_paid"] == 2000000
    assert item_a["total_debt"] == 13000000
    assert item_a["debt_0_30"] == 8000000
    assert item_a["debt_31_60"] == 5000000
    assert item_a["debt_61_90"] == 0
    assert item_a["debt_over_90"] == 0
    assert item_a["overdue_debt"] == 5000000
    assert item_a["is_overdue"] is True
    assert item_a["credit_limit"] == 50000000

    # Kiểm tra số liệu Đại lý B
    assert item_b["total_receivable"] == 50000000
    assert item_b["total_paid"] == 5000000
    assert item_b["total_debt"] == 45000000
    assert item_b["debt_0_30"] == 0
    assert item_b["debt_31_60"] == 0
    assert item_b["debt_61_90"] == 15000000
    assert item_b["debt_over_90"] == 30000000
    assert item_b["overdue_debt"] == 45000000
    assert item_b["is_overdue"] is True

    # Kiểm tra dòng TỔNG CỘNG chân bảng (SCRUM-721)
    summary = data["summary"]
    assert summary["total_customers"] == 2
    assert summary["total_receivable"] == 65000000
    assert summary["total_paid"] == 7000000
    assert summary["total_debt"] == 58000000
    assert summary["total_debt_0_30"] == 8000000
    assert summary["total_debt_31_60"] == 5000000
    assert summary["total_debt_61_90"] == 15000000
    assert summary["total_debt_over_90"] == 30000000
    assert summary["total_overdue_debt"] == 50000000
    assert summary["overdue_customers_count"] == 2


def test_03_filters_scrum719():
    """
    SCRUM-719: Kiểm thử các bộ lọc (search, region, customer_group, aging_group, only_overdue).
    """
    client = TestClient(app)
    headers = auth_header(1, "admin_debt", UserRole.ADMIN.value)

    # 1. Lọc theo search DL-001
    res1 = client.get(
        "/api/v1/agency-debt-reports?search=Tuấn Phương&as_of_date=2026-10-15",
        headers=headers,
    )
    assert res1.status_code == 200
    assert res1.json()["total"] == 1
    assert res1.json()["items"][0]["customer_code"] == "DL-001"

    # 2. Lọc theo region "Miền Nam"
    res2 = client.get(
        "/api/v1/agency-debt-reports?region=Miền Nam&as_of_date=2026-10-15",
        headers=headers,
    )
    assert res2.status_code == 200
    assert res2.json()["total"] == 1
    assert res2.json()["items"][0]["customer_code"] == "DL-002"

    # 3. Lọc theo aging_group "over_90" -> Chỉ DL-002 có nợ trên 90 ngày
    res3 = client.get(
        "/api/v1/agency-debt-reports?aging_group=over_90&as_of_date=2026-10-15",
        headers=headers,
    )
    assert res3.status_code == 200
    assert res3.json()["total"] == 1
    assert res3.json()["items"][0]["customer_code"] == "DL-002"

    # 4. Lọc theo aging_group "0_30" -> Chỉ DL-001 có nợ nhóm 0-30 ngày
    res4 = client.get(
        "/api/v1/agency-debt-reports?aging_group=0_30&as_of_date=2026-10-15",
        headers=headers,
    )
    assert res4.status_code == 200
    assert res4.json()["total"] == 1
    assert res4.json()["items"][0]["customer_code"] == "DL-001"


def test_04_export_excel_scrum722():
    """
    SCRUM-722: Phát triển chức năng xuất Excel cho báo cáo sổ công nợ theo đại lý.
    Kiểm tra workbook hợp lệ, tiêu đề, các cột tuổi nợ, định dạng số và dòng tổng cộng.
    """
    client = TestClient(app)
    headers = auth_header(2, "accountant_debt", UserRole.ACCOUNTANT.value)

    res = client.get(
        "/api/v1/agency-debt-reports/export-excel?as_of_date=2026-10-15",
        headers=headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in res.headers["content-disposition"]

    # Đọc lại nội dung excel bằng openpyxl
    file_bytes = io.BytesIO(res.content)
    wb = openpyxl.load_workbook(file_bytes, data_only=True)
    assert "So_Cong_No_Dai_Ly" in wb.sheetnames
    ws = wb["So_Cong_No_Dai_Ly"]

    # Kiểm tra dòng tiêu đề
    assert "BÁO CÁO SỔ CÔNG NỢ" in ws["A1"].value
    assert "2026" in ws["A2"].value

    # Kiểm tra header row (row 5)
    headers_list = [ws.cell(5, c).value for c in range(1, 17)]
    assert "Mã đại lý" in headers_list
    assert "Còn nợ (VNĐ)" in headers_list
    assert "Tuổi nợ 0–30 ngày" in headers_list
    assert "Tuổi nợ >90 ngày" in headers_list
    assert "Nợ quá hạn (VNĐ)" in headers_list

    # Kiểm tra dòng dữ liệu (row 6, 7)
    row6_code = ws.cell(6, 2).value
    assert row6_code in ("DL-001", "DL-002")

    # Kiểm tra dòng tổng cộng chân bảng (row 8)
    footer_label = ws.cell(8, 1).value
    assert "TỔNG CỘNG" in str(footer_label)
    # Cột 9 là Còn nợ tổng cộng = 58,000,000
    assert ws.cell(8, 9).value == 58000000


def test_05_agency_invoices_drilldown():
    """
    Kiểm thử API drill-down xem danh sách hóa đơn cấu thành công nợ của 1 đại lý.
    """
    client = TestClient(app)
    headers = auth_header(2, "accountant_debt", UserRole.ACCOUNTANT.value)

    res = client.get(
        "/api/v1/agency-debt-reports/CUS-A/invoices?as_of_date=2026-10-15",
        headers=headers,
    )
    assert res.status_code == 200
    data = res.json()
    assert data["customer"]["customer_code"] == "DL-001"
    assert len(data["invoices"]) == 2

    inv1 = next(i for i in data["invoices"] if i["invoice_id"] == "HD-A1")
    assert inv1["remaining_amount"] == 8000000
    assert inv1["aging_bucket"] == "0_30"
    assert inv1["is_overdue"] is False

    inv2 = next(i for i in data["invoices"] if i["invoice_id"] == "HD-A2")
    assert inv2["remaining_amount"] == 5000000
    assert inv2["aging_bucket"] == "31_60"
    assert inv2["is_overdue"] is True
    assert inv2["days_overdue"] == 10


def test_06_customer_and_sales_rep_scope():
    """
    Kiểm thử phân quyền truy cập:
    - Tài khoản vai trò Customer bị chặn tuyệt đối khỏi phân hệ kế toán nội bộ (403 Forbidden).
    - Nhân viên kinh doanh (Sales Rep) chỉ xem được đại lý mình được gán phụ trách.
    """
    client = TestClient(app)

    # 1. Customer gọi vào route kế toán nội bộ -> Bị chặn 403 Forbidden theo quy tắc phân quyền hệ thống
    cust_headers = auth_header(4, "customer_user", UserRole.CUSTOMER.value)
    res_cust = client.get(
        "/api/v1/agency-debt-reports?as_of_date=2026-10-15",
        headers=cust_headers,
    )
    assert res_cust.status_code == 403
    assert "quản trị nội bộ" in res_cust.json()["detail"]

    # 2. Sales Rep (gán phụ trách DL-001) gọi -> Chỉ thấy DL-001 (CUS-A)
    sales_headers = auth_header(3, "sales_rep_debt", UserRole.SALES_REP.value)
    res_sales = client.get(
        "/api/v1/agency-debt-reports?as_of_date=2026-10-15",
        headers=sales_headers,
    )
    assert res_sales.status_code == 200
    sales_data = res_sales.json()
    assert sales_data["total"] == 1
    assert sales_data["items"][0]["customer_code"] == "DL-001"
