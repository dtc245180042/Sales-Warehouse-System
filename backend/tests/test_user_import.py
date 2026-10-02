import io
import pytest
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, lay_phien_db
from app.core.security import bam_mat_khau, tao_token_truy_cap
from app.models.auth import User, UserRole
from app.services.seed_service import seed_all
from main import app

# Database SQLite in-memory tách biệt cho test suite SC-209
TEST_DB_URL = "sqlite:///:memory:"
test_engine = create_engine(
    TEST_DB_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_lay_phien_db():
    phien_db = TestingSessionLocal()
    try:
        yield phien_db
    finally:
        phien_db.close()


@pytest.fixture(autouse=True)
def setup_test_db():
    """Tạo schema và dữ liệu mẫu sạch trước mỗi test."""
    Base.metadata.create_all(bind=test_engine)
    app.dependency_overrides[lay_phien_db] = override_lay_phien_db
    db = TestingSessionLocal()
    seed_all(db)

    # Tạo sẵn một Admin để test
    admin_user = User(
        username="admin_importer",
        email="admin_importer@company.vn",
        hashed_password=bam_mat_khau("AdminPass@123"),
        role=UserRole.ADMIN.value,
        is_active=True,
        token_version=1,
    )
    # Tạo sẵn một User thường để test DB duplicate
    existing_user = User(
        username="existing_sales",
        email="existing_sales@company.vn",
        hashed_password=bam_mat_khau("Existing@123"),
        role=UserRole.SALES_REP.value,
        is_active=True,
        token_version=1,
    )
    db.add(admin_user)
    db.add(existing_user)
    db.commit()
    db.close()

    yield

    Base.metadata.drop_all(bind=test_engine)
    app.dependency_overrides.clear()


def lay_token_admin() -> str:
    db = TestingSessionLocal()
    admin = db.query(User).filter(User.username == "admin_importer").first()
    token = tao_token_truy_cap({"user_id": admin.id, "sub": str(admin.id), "token_version": admin.token_version})
    db.close()
    return token


def lay_token_user_thuong() -> str:
    db = TestingSessionLocal()
    user = db.query(User).filter(User.username == "existing_sales").first()
    token = tao_token_truy_cap({"user_id": user.id, "sub": str(user.id), "token_version": user.token_version})
    db.close()
    return token


def tao_file_excel(rows_data: list, headers: list = None) -> bytes:
    """Helper tạo file Excel nhị phân cho test."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DanhSachNguoiDung"

    if headers is None:
        headers = [
            "username", "email", "full_name", "role",
            "phone_number", "assigned_warehouse", "password"
        ]
    ws.append(headers)
    for row in rows_data:
        ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ----------------- TEST CASES CHO SC-209 -----------------

def test_download_template_admin_thanh_cong():
    """Kiểm tra API GET /api/v1/user-imports/template cho phép Admin tải file mẫu hợp lệ."""
    client = TestClient(app)
    token = lay_token_admin()
    res = client.get(
        "/api/v1/user-imports/template",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res.headers["content-type"]
    assert "attachment; filename=\"user_import_template.xlsx\"" in res.headers.get("content-disposition", "")

    # Đọc lại file xem có đúng 2 sheet và có header tiếng Việt không
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "DanhSachNguoiDung" in wb.sheetnames
    assert "HuongDan" in wb.sheetnames
    ws = wb["DanhSachNguoiDung"]
    assert ws.cell(row=1, column=1).value == "Tên đăng nhập"
    assert ws.cell(row=1, column=2).value == "Email"


def test_preview_va_import_tu_dong_lay_email_khi_de_trong_username():
    """Kiểm tra khi để trống username thì hệ thống tự động lấy Email làm username."""
    client = TestClient(app)
    token = lay_token_admin()

    headers_vn = [
        "Tên đăng nhập", "Email", "Họ và tên", "Vai trò",
        "Số điện thoại", "Kho / Địa bàn phụ trách", "Mật khẩu"
    ]
    rows = [
        # Dòng để trống username
        ("", "sales_auto@company.vn", "Nguyễn Tự Động", "Sales Rep", "0911223344", "Địa bàn Miền Bắc", "Pass@123"),
    ]

    excel_bytes = tao_file_excel(rows, headers=headers_vn)
    files = {"file": ("test_auto.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    # 1. Preview
    res_prev = client.post(
        "/api/v1/user-imports/preview",
        headers={"Authorization": f"Bearer {token}"},
        files=files,
    )
    assert res_prev.status_code == 200
    data = res_prev.json()
    assert data["valid_count"] == 1
    assert data["rows"][0]["is_valid"] is True
    assert data["rows"][0]["username"] == "sales_auto@company.vn"

    # 2. Execute import
    files = {"file": ("test_auto.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res_exec = client.post(
        "/api/v1/user-imports/execute",
        headers={"Authorization": f"Bearer {token}"},
        files=files,
    )
    assert res_exec.status_code == 200
    summary = res_exec.json()
    assert summary["success_count"] == 1
    assert summary["created_users"][0]["username"] == "sales_auto@company.vn"

    # 3. Check DB
    db = TestingSessionLocal()
    u = db.query(User).filter(User.email == "sales_auto@company.vn").first()
    assert u is not None
    assert u.username == "sales_auto@company.vn"
    db.close()


def test_download_template_khong_phai_admin_bi_tu_choi():
    """Kiểm tra người dùng không phải Admin không thể tải tệp template (403 Forbidden)."""
    client = TestClient(app)
    token = lay_token_user_thuong()
    res = client.get(
        "/api/v1/user-imports/template",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res.status_code == 403


def test_preview_dinh_dang_file_khong_hop_le():
    """Tải lên file không phải Excel (.txt, .pdf) sẽ bị từ chối 400 Bad Request."""
    client = TestClient(app)
    token = lay_token_admin()
    files = {"file": ("test.txt", b"dummy content", "text/plain")}
    res = client.post(
        "/api/v1/user-imports/preview",
        headers={"Authorization": f"Bearer {token}"},
        files=files,
    )
    assert res.status_code == 400
    assert "Định dạng tệp không được hỗ trợ" in res.json()["detail"]


def test_preview_va_bao_loi_tung_dong():
    """Kiểm tra preview phân loại chính xác các dòng hợp lệ và báo lỗi chi tiết từng dòng lỗi."""
    client = TestClient(app)
    token = lay_token_admin()

    rows = [
        # Dòng 2: Hợp lệ hoàn toàn
        ("sales_01", "sales01@company.vn", "Trần Văn Một", "Sales Rep", "0901234567", "", "Pass@123"),
        # Dòng 3: Trùng username với DB (existing_sales)
        ("existing_sales", "new_email@company.vn", "Trần Văn Hai", "Sales Rep", "", "", ""),
        # Dòng 4: Sai định dạng email
        ("sales_03", "invalid-email-format", "Trần Văn Ba", "Sales Rep", "", "", ""),
        # Dòng 5: Vai trò Kho nhưng thiếu kho phụ trách (SCRUM-206)
        ("wh_staff_01", "wh01@company.vn", "Nguyễn Kho", "Warehouse", "0912345678", "", ""),
        # Dòng 6: Vai trò không hợp lệ
        ("unknown_01", "unknown@company.vn", "Người Lạ", "SuperHero", "", "", ""),
        # Dòng 7: Trùng lặp nội bộ trong cùng file (trùng email với dòng 2)
        ("sales_07", "sales01@company.vn", "Trần Văn Bảy", "Sales Rep", "", "", ""),
    ]

    excel_bytes = tao_file_excel(rows)
    files = {"file": ("users.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    res = client.post(
        "/api/v1/user-imports/preview",
        headers={"Authorization": f"Bearer {token}"},
        files=files,
    )
    assert res.status_code == 200
    data = res.json()

    assert data["total_rows"] == 6
    assert data["valid_count"] == 1
    assert data["invalid_count"] == 5

    # Dòng 2 hợp lệ
    row_2 = next(r for r in data["rows"] if r["row_index"] == 2)
    assert row_2["is_valid"] is True
    assert len(row_2["errors"]) == 0

    # Dòng 3 báo lỗi trùng DB
    row_3 = next(r for r in data["rows"] if r["row_index"] == 3)
    assert row_3["is_valid"] is False
    assert any("đã tồn tại trong hệ thống" in e for e in row_3["errors"])

    # Dòng 4 báo lỗi email
    row_4 = next(r for r in data["rows"] if r["row_index"] == 4)
    assert row_4["is_valid"] is False
    assert any("Định dạng email" in e for e in row_4["errors"])

    # Dòng 5 báo lỗi thiếu kho
    row_5 = next(r for r in data["rows"] if r["row_index"] == 5)
    assert row_5["is_valid"] is False
    assert any("Kho phụ trách" in e for e in row_5["errors"])

    # Dòng 6 báo lỗi role
    row_6 = next(r for r in data["rows"] if r["row_index"] == 6)
    assert row_6["is_valid"] is False
    assert any("không hợp lệ" in e for e in row_6["errors"])

    # Dòng 7 báo lỗi trùng với dòng 2 trong file
    row_7 = next(r for r in data["rows"] if r["row_index"] == 7)
    assert row_7["is_valid"] is False
    assert any("trùng với dòng 2" in e for e in row_7["errors"])


def test_execute_partial_import_va_bao_cao_tong_ket():
    """Kiểm tra cơ chế partial import: Lưu các dòng hợp lệ, bỏ qua các dòng lỗi và xuất báo cáo tổng kết."""
    client = TestClient(app)
    token = lay_token_admin()

    rows = [
        # Dòng hợp lệ 1
        ("nv_sales_10", "sales10@company.vn", "Nguyễn Số Mười", "Sales Rep", "0911223344", "", "Pass@10"),
        # Dòng lỗi (trùng username với DB)
        ("existing_sales", "another_existing@company.vn", "Trùng Tên", "Customer", "", "", ""),
        # Dòng hợp lệ 2 (vai trò Kho có kho phụ trách)
        ("nv_kho_11", "kho11@company.vn", "Lê Văn Kho", "Warehouse", "0988776655", "Kho Hà Nội", ""),
        # Dòng lỗi (sai email)
        ("nv_sales_12", "bad-email", "Hỏng Email", "Sales Rep", "", "", ""),
    ]

    excel_bytes = tao_file_excel(rows)
    files = {"file": ("batch_users.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    res = client.post(
        "/api/v1/user-imports/execute",
        headers={"Authorization": f"Bearer {token}"},
        files=files,
    )
    assert res.status_code == 200
    summary = res.json()

    assert summary["total_processed"] == 4
    assert summary["success_count"] == 2
    assert summary["failed_count"] == 2
    assert len(summary["created_users"]) == 2
    assert len(summary["row_errors"]) == 2

    # Kiểm tra người dùng đã thực sự được lưu vào DB
    db = TestingSessionLocal()
    u1 = db.query(User).filter(User.username == "nv_sales_10").first()
    assert u1 is not None
    assert u1.email == "sales10@company.vn"
    assert u1.role == "Sales Rep"
    assert u1.must_change_password is False

    u2 = db.query(User).filter(User.username == "nv_kho_11").first()
    assert u2 is not None
    assert u2.assigned_warehouse == "Kho Hà Nội"
    # Do để trống password nên hệ thống tự cấp mật khẩu tạm
    assert u2.must_change_password is True
    db.close()
